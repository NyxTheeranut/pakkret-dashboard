#!/usr/bin/env python3
"""
Pushes Pak Kret Territory Explorer's RTR + 7-11 data into the Google Sheet the
dashboard's Apps Script backend serves from -- no browser, no file picker, no
Google sign-in. Run this whenever the source xlsx files change.

This is a line-by-line port of the exact same ingestion/aggregation pipeline
the dashboard's own JavaScript runs (Dashboards/Pak Kret Territory
Explorer.html), NOT a re-derivation from scratch. Every function below is
named after, and mirrors, its JS counterpart -- when that file's parsing logic
changes, this needs the same change made here, by hand, side by side with the
JS. There is no automated way to keep the two in sync; that's the accepted
tradeoff for not needing a browser open to publish (see the plan this was
built from for the full reasoning).

What this does NOT do: recompute %Achieve, MOM%, or GA (LMTD) for display.
Those are only ever derived client-side, on every load, straight from the raw
ga/target/daily numbers this script writes (see the dashboard's
finalizeRtrMonths/sevRowData) -- so this script only needs to get the RAW
numbers and each store's TARGET right, not any of the derived display math.

Data written, per month, mirrors buildRtrMonthTables/buildSevMonthTables
exactly (same tab-name convention, same column order) -- see PUBLISH_URL
setup below.

── SETUP (one-time) ────────────────────────────────────────────────────────
 1. In the "PakKret Data" Sheet's Apps Script project (Extensions -> Apps
    Script), paste in the current "Sheets Sync - Apps Script Code.gs" from
    this folder, which now includes the publishTerritory action.
 2. Project Settings -> Script Properties -> add TERRITORY_SYNC_SECRET, the
    exact value in Config/territory_sync_secret.txt.
 3. Deploy -> Manage deployments -> edit the existing Web app deployment ->
    New version -> Deploy.
 4. Run this script. It reads:
      Territory Explorer/Source Data/RTR/*.xlsx  (the Commit Target GA file)
      Territory Explorer/Source Data/7-11/*.xlsx (Q1/Q2/Q3 sales, roster,
        target/commit files -- same multi-file picker convention the
        dashboard's 7-11 tab already uses; anything that doesn't look like
        sales/roster/target data is skipped with a warning, not a crash)
    and pushes everything through the new publishTerritory Apps Script
    action, gated by the shared secret above.
"""
import json
import re
import sys
import time
import urllib.error
import urllib.request
import warnings
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl

# These are cosmetic complaints about the xlsx's own formatting (a print area, a
# conditional-formatting rule, a sparkline group openpyxl doesn't preserve on
# read) -- never about the data this script actually reads. Left enabled they
# interrupt the progress output below at unpredictable points with several
# lines of noise the terminal has no use for.
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Keep this in sync with VIEWER_SCRIPT_URL in Dashboards/Pak Kret Territory
# Explorer.html -- same deployment, same URL, this just adds a new action to it.
SYNC_URL = "https://script.google.com/macros/s/AKfycbzjNSiIzcG2d0o1VcGX-kVoUtj3HBsnyDCMrO1qBimiIBUN5N3qLvQ4rkio30yVkxmv/exec"

# Matches SYNC.KEEP_MONTHS in the dashboard JS.
KEEP_MONTHS = 6

DASHBOARD_DIR = Path(__file__).resolve().parent.parent
TERRITORY_DIR = DASHBOARD_DIR / "Territory Explorer"
RTR_DIR = TERRITORY_DIR / "Source Data" / "RTR"
SEV_DIR = TERRITORY_DIR / "Source Data" / "7-11"
SYNC_SECRET_FILE = DASHBOARD_DIR / "Config" / "territory_sync_secret.txt"


def _find_latest(folder, pattern):
    # "~$..." lock files Excel creates while the real file is open sort FIRST here (reverse=True,
    # and "~" is a high-value character), so without this filter, having the source workbook open
    # in Excel makes this pick the lock file instead of the real one -- exactly what just happened
    # (BadZipFile: a "~$..." file isn't a real xlsx, just a few bytes of open-file metadata).
    # SEV_PATHS below already filters these; this one didn't.
    candidates = sorted(
        (p for p in folder.glob(pattern) if not p.name.startswith("~$")),
        key=lambda p: p.name,
        reverse=True,
    )
    return candidates[0] if candidates else None


def _find_rtr_workbook(folder):
    # Filename-sort alone (_find_latest's usual trick) assumed exactly one .xlsx ever lives in
    # this folder. That stopped being true the day a second, differently-shaped RTR workbook
    # ("RTR_Commit Routing...") landed alongside the real target/mapping workbook -- it happens
    # to sort first by name, but has NONE of the sheets this pipeline depends on (no "RTR Target
    # GA" mapping sheet, no "KPI Set Up"). process_rtr_workbook() only hard-fails if a candidate
    # has zero month sheets; a missing mapping sheet degrades silently instead (code_to_canonical
    # ends up empty, so every store looks unmapped) -- exactly the kind of wrong-but-not-crashing
    # result that could get published without anyone noticing. So: actually open each candidate
    # (sheet names only, cheap) and skip any that lack a real mapping sheet, newest-by-name first.
    candidates = sorted(
        (p for p in folder.glob("*.xlsx") if not p.name.startswith("~$")),
        key=lambda p: p.name,
        reverse=True,
    )
    skipped = []
    for p in candidates:
        try:
            wb = openpyxl.load_workbook(p, read_only=True)
            names = wb.sheetnames
            wb.close()
        except Exception as e:
            skipped.append(f"{p.name} (could not open: {e})")
            continue
        if any(MAP_SHEET_RE.search(n) for n in names):
            if skipped:
                print(f"  (skipped, missing a 'RTR Target GA' mapping sheet: "
                      f"{', '.join(skipped)})")
            return p
        skipped.append(p.name)
    if skipped:
        raise SystemExit(
            f"No usable RTR workbook in {folder} -- none of these have a sheet matching "
            f"'RTR Target GA': {', '.join(skipped)}"
        )
    return None


def _find_routing_workbook(folder, exclude=None):
    # The "routing" workbook carries data the primary workbook above doesn't have: a live
    # per-store target for whichever month it's currently tracking, and each store's COM (the
    # individual field rep one level below CM). Matched purely by its own "RTR_Route" sheet,
    # never by filename -- the filename's quarter/year is expected to keep changing. Optional:
    # returns None (not an error) if nothing matches, since every caller treats this as extra
    # data a publish can proceed without.
    candidates = sorted(
        (p for p in folder.glob("*.xlsx") if not p.name.startswith("~$") and p != exclude),
        key=lambda p: p.name,
        reverse=True,
    )
    for p in candidates:
        try:
            wb = openpyxl.load_workbook(p, read_only=True)
            names = wb.sheetnames
            wb.close()
        except Exception:
            continue
        if any(ROUTING_ROUTE_SHEET_RE.match(n.strip()) for n in names):
            return p
    return None


# ═════════════════════════════════════════════════════════════════════════
# SHARED CONSTANTS -- ported verbatim from the top of the dashboard's <script>
# ═════════════════════════════════════════════════════════════════════════
TARGET_CLUSTER_KEYWORD = "pak kret"
SEV_NAME_RE = re.compile(r"7-11|7-eleven", re.I)
MONTH_RE = re.compile(r"^(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)", re.I)
MAP_SHEET_RE = re.compile(r"rtr target ga", re.I)
KPI_SETUP_RE = re.compile(r"kpi\s*set\s*up", re.I)
# Matches the "routing" workbook's own distinctive sheet (currently "RTR_Commit Routing
# Q42026.xlsx", presumably "...Q12027.xlsx" etc. next quarter) -- matched by sheet name, NOT
# filename, since the quarter/year in the filename is expected to keep changing.
ROUTING_ROUTE_SHEET_RE = re.compile(r"^rtr_route$", re.I)
# Per-CM committed targets (GA + M1, Mass+Migrant/Tourist splits, and Ach M1) for months the
# primary workbook's KPI Set Up doesn't cover yet -- same block layout parse_kpi_setup_targets
# already reads, just on this sheet instead.
PERF_CM_SHEET_RE = re.compile(r"^performance\s*cm", re.I)
PERF_COM_SHEET_RE = re.compile(r"^performance\s*com\b", re.I)
SEV_TARGET_SHEET_RE = re.compile(r"^tg\s", re.I)
SEV_CM_TARGET_SHEET_RE = re.compile(r"^tg\s*cm\s*$", re.I)

REQUIRED_COLS = ["PARTNER_CODE", "PARTNER_NAME", "GA", "GROUP_SIM", "D_CLUSTER", "TM_KEY_DAY"]
MAP_REQUIRED_NORM = ["retailer name", "cluster", "cm"]
TIER_COLS = [
    "AP1D_1-49", "AP1D_50-99", "AP1D_100-119", "AP1D_120-149",
    "AP1D_150-199", "AP1D_200-249", "AP1D_250-299", "AP1D_300UP",
]
TIER_LABELS = ["1-49", "50-99", "100-119", "120-149", "150-199", "200-249", "250-299", "300+"]
MONTH_ABBR = {"jan": 0, "feb": 1, "mar": 2, "apr": 3, "may": 4, "jun": 5,
              "jul": 6, "aug": 7, "sep": 8, "oct": 9, "nov": 10, "dec": 11}
MONTH_LABEL = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
UNMAPPED_CM = "ไม่ระบุ CM"
UNMAPPED_AE = "ไม่ระบุ AE"

# The RTR mapping sheet's own "อำเภอ" column is in Thai, but every monthly sheet's
# DISTRICT_EN column (the source for an ACTIVE store's district -- see aggregate_rtr_sheet)
# is in English -- normalized here so a dormant store (fill_dormant_roster_stores, which has
# no monthly-sheet row of its own to pull DISTRICT_EN from and must fall back to this Thai
# mapping-sheet column) gets the SAME label an active store in the same district would,
# instead of silently splitting one district into two separate filter pills (one per
# language) the way it did before this normalization existed. Limited to the three อำเภอ
# this whole territory is ever scoped to (TARGET_CLUSTER_KEYWORD) -- not a general Thai/
# English gazetteer.
DISTRICT_TH_TO_EN = {
    "ปากเกร็ด": "Pak Kret",
    "บางบัวทอง": "Bang Bua Thong",
    "ไทรน้อย": "Sai Noi",
}

# Store codes confirmed (with the user) to not belong in this dataset at all -- same
# treatment as an out-of-cluster CM (see the known_cms scoping in process_sev_files):
# excluded entirely, never counted under any CM/AE, not reassigned. 71100297 ("7-11
# แจ้งวัฒนะ 28(NC)") only ever shows up in Target 7-eleven Aug'26_BMA West.xlsx's "Data
# Aug" sheet with Status="Close" -- a closed store lingering in one source file.
EXCLUDED_STORE_CODES = {"71100297"}

# SIM_TYPE values that are promotional/device-bundle activations, not a priced SIM plan
# -- sim_exact_price_band already correctly leaves these unbanded (no AP49/99/199/>199
# tier fits them, since there's no price to parse out of the name), but by request they
# also shouldn't count toward GA/revenue totals at all. Filtered out here, before
# publishing -- the source xlsx files themselves are left untouched.
EXCLUDED_SIM_TYPES = {"7-11 DEVICE BUNDLING"}

SEV_REQUIRED_HEADERS = [
    "TM_KEY_DAY", "TDS_PROVINCE", "DISTRICT_EN", "PARTNER_CODE",
    "PARTNER_NAME", "GROUP_SIM", "Sum of GA",
]
SEV_TARGET_REQUIRED_NORM = ["partner_code", "cm"]


# ═════════════════════════════════════════════════════════════════════════
# SHARED HELPERS -- same names as the JS (num, normHeader, inTargetCluster, ...)
# ═════════════════════════════════════════════════════════════════════════
def get(row, i):
    """row[i], but tolerant of short/ragged rows the way JS's undefined-on-any-
    out-of-bounds-index always was -- openpyxl rows aren't guaranteed padded."""
    return row[i] if row is not None and 0 <= i < len(row) else None


def num(v):
    if isinstance(v, bool):
        return 0
    if isinstance(v, (int, float)):
        return v
    if v is None:
        return 0
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return 0


def norm_header(h):
    return re.sub(r"\s+", " ", "" if h is None else str(h)).strip().lower()


def in_target_cluster(v):
    return v is not None and TARGET_CLUSTER_KEYWORD in str(v).lower()


def day_of_month(day_key):
    s = str(day_key)
    try:
        return int(s[6:8])
    except ValueError:
        return 0


def parse_sheet_label(name):
    m = re.match(r"^([a-zA-Z]{3})[a-zA-Z]*\.?(\d{2})?", name.strip())
    if not m:
        return name
    abbr = m.group(1).lower()
    if abbr not in MONTH_ABBR:
        return name
    label = MONTH_LABEL[MONTH_ABBR[abbr]]
    if m.group(2):
        label += " 20" + m.group(2)
    return label


def parse_sheet_date(name):
    m = re.match(r"^([a-zA-Z]{3})[a-zA-Z]*\.?(\d{2})?", name.strip())
    if not m:
        return None
    abbr = m.group(1).lower()
    if abbr not in MONTH_ABBR:
        return None
    return {"monthIndex": MONTH_ABBR[abbr], "year": (2000 + int(m.group(2))) if m.group(2) else None}


def parse_month_from_filename(name):
    """Extracts a 'YYYYMM' month key from an arbitrary filename, e.g. "Target 7-eleven
    Aug'26_BMA West.xlsx" -> "202608". Unlike parse_sheet_date, the month doesn't have to
    be at the very start of the string -- a target/commit file's month is embedded
    somewhere in a longer filename, not the whole name itself."""
    m = re.search(
        # (?!\d) rather than \b at the end: \b doesn't fire between a digit and "_" (both are
        # \w), and filenames like "Aug'26_BMA West.xlsx" have exactly that -- a plain \b here
        # silently fails to match on real filenames, falling back to the wrong month.
        r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*['\s]?(\d{2})(?!\d)",
        name,
        re.I,
    )
    if not m:
        return None
    abbr = m.group(1).lower()
    if abbr not in MONTH_ABBR:
        return None
    return f"{2000 + int(m.group(2))}{MONTH_ABBR[abbr] + 1:02d}"


def sort_months_chronologically(months):
    years = [d["year"] for d in (parse_sheet_date(m["sheetName"]) for m in months) if d and d["year"]]
    inferred_year = Counter(years).most_common(1)[0][0] if years else date.today().year
    for m in months:
        d = parse_sheet_date(m["sheetName"]) or {"monthIndex": 0, "year": inferred_year}
        m["_sortKey"] = (d["year"] or inferred_year) * 100 + d["monthIndex"]
    months.sort(key=lambda m: m["_sortKey"])
    return months


def find_header_row(rows2d, required_norm, max_scan=20):
    for i in range(min(max_scan, len(rows2d))):
        row = rows2d[i]
        if not row:
            continue
        normed = [norm_header(c) for c in row]
        if all(rn in normed for rn in required_norm):
            return i
    return -1


def sheet_to_rows(ws, max_row=None):
    """Same shape as XLSX.utils.sheet_to_json(ws, {header:1, raw:true}) --
    a plain list of row-lists, values only, row 0 is whatever's actually in
    the sheet's first row (not necessarily a real header)."""
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if max_row is not None and i >= max_row:
            break
        rows.append(list(row))
    return rows


# ═════════════════════════════════════════════════════════════════════════
# RTR PIPELINE
# ═════════════════════════════════════════════════════════════════════════

def parse_target_blocks(header):
    """Finds each 'GA by RSR <Month> Target' anchor column, then locates that
    month's Actual/M1/%Achieve sub-columns within the window up to the next
    anchor. Positional, not name-keyed, because the sheet repeats
    near-identical column names once per month block. Ported from
    parseTargetBlocks in the dashboard JS."""
    normed = [norm_header(h) for h in header]
    anchors = []
    for i, val in enumerate(normed):
        if val.startswith("ga by") and "target" in val:
            mo = next((abbr for abbr in MONTH_ABBR if abbr in val), None)
            if mo:
                anchors.append({"idx": i, "month": mo})

    blocks = []
    for a_i, anchor in enumerate(anchors):
        start = anchor["idx"]
        end = anchors[a_i + 1]["idx"] if a_i + 1 < len(anchors) else len(header)
        win = normed[start:end]

        def find_in(pred, win=win, start=start):
            for k, c in enumerate(win):
                if pred(c):
                    return start + k
            return -1

        blocks.append({
            "month": anchor["month"],
            "targetIdx": start,
            "actualIdx": find_in(lambda c: c.startswith("ga") and ("actual" in c or "mtd" in c)),
            "m1Idx": find_in(lambda c: c.startswith("m1")),
            "achieveIdx": find_in(lambda c: c == "% achieve"),
        })
    return blocks


def parse_kpi_setup_targets(rows2d, name_col="cm"):
    """The 'KPI Set Up' sheet: a leadership-committed target per CM, laid out
    as a two-row header repeated once per month, at three granularities (PBH/
    CM/RSR) -- only the CM-level block, and only when its own 'Total' row's
    'All RTR' cell has a real number (the sheet's own sanity check), is
    trusted. Ported from parseKpiSetUpTargets. Returns {monthAbbr: {cmName:
    {total, massMigrant, tourist, m1Total, m1MassMigrant, m1Tourist,
    achM1Total, achM1MassMigrant, achM1Tourist}}}, or None if nothing usable
    was found.

    massMigrant/tourist: the same "Target GA (<Month>)" block also carries a
    2-way segment split immediately before its own "Total GA" column ("Mass +
    Migrant" bundled as one number, "Tourist" separate -- the sheet never
    splits Mass from Migrant individually). Captured here (not previously)
    so the dashboard can show a segment-scoped Target when the Mass/Migrant/
    Tourist trend-chart toggles are in a state that maps to a real number --
    see rtrTargetForSegToggle in the dashboard JS.

    m1Total/m1MassMigrant/m1Tourist: a SEPARATE "Target M1 (<Month>)" anchor
    sits later in the exact same header row as "Target GA (<Month>)" (every
    month, every granularity -- confirmed against the real file), same
    3-column Mass+Migrant/Tourist/Total shape, just never read before this.
    Found independently of the GA block so a CM present in one but not the
    other (in principle) still gets whichever side is real -- not previously
    read at all, this was the actual gap behind "RTR has no M1 target"."""
    results = {}
    for g_idx, g_row in enumerate(rows2d):
        if not g_row:
            continue
        anchors = [{"idx": i, "text": norm_header(v)} for i, v in enumerate(g_row) if norm_header(v)]
        target_anchor = next((a for a in anchors if a["text"].startswith("target ga")), None)
        if not target_anchor:
            continue
        if g_idx + 1 >= len(rows2d) or not rows2d[g_idx + 1]:
            continue
        s_row = rows2d[g_idx + 1]
        s_norm = [norm_header(c) for c in s_row]
        # name_col="com_name" reads the routing workbook's "Performance COM" sheet, which has
        # this exact block layout, just one row per COM instead of per CM.
        if name_col not in s_norm:
            continue  # PBH-only or RSR-level block -- keep scanning past it
        cm_idx = s_norm.index(name_col)
        all_rtr_idx = s_norm.index("all rtr") if "all rtr" in s_norm else -1

        def block_bounds(anchor_idx):
            end = len(g_row)
            for a2 in anchors:
                if a2["idx"] > anchor_idx:
                    end = a2["idx"]
                    break
            return end

        # First occurrence wins for each, same as the original single-column scan --
        # matters because "Tourist" (unlike "Total GA"/"Mass + Migrant") can legitimately
        # repeat later in the row as its own granularity's column in some blocks.
        def scan_block(anchor_idx, total_label):
            end = block_bounds(anchor_idx)
            total_idx, mm_idx, tr_idx = -1, -1, -1
            for c in range(anchor_idx, end):
                if c >= len(s_norm):
                    continue
                if total_idx == -1 and s_norm[c] == total_label:
                    total_idx = c
                elif mm_idx == -1 and s_norm[c] == "mass + migrant":
                    mm_idx = c
                elif tr_idx == -1 and s_norm[c] == "tourist":
                    tr_idx = c
            return total_idx, mm_idx, tr_idx

        total_ga_idx, mass_migrant_idx, tourist_idx = scan_block(target_anchor["idx"], "total ga")
        if total_ga_idx == -1:
            continue

        month_abbr = next((abbr for abbr in MONTH_ABBR if abbr in target_anchor["text"]), None)
        if not month_abbr or month_abbr in results:
            continue  # already have this month's CM-level block

        m1_anchor = next(
            (a for a in anchors if a["text"].startswith("target m1") and month_abbr in a["text"]),
            None,
        )
        total_m1_idx, mm_m1_idx, tr_m1_idx = (
            scan_block(m1_anchor["idx"], "total m1") if m1_anchor else (-1, -1, -1)
        )
        # "Ach M1 (<Month>)" is its own anchor immediately after the Target M1 block (same row,
        # same 3-column Mass+Migrant/Tourist/Total M1 shape) -- the CM's actual M1 as this sheet
        # itself records it, more authoritative than summing each store's own "M1 <Month>" column
        # from the mapping sheet (a different source entirely, by request: use this one instead).
        ach_m1_anchor = next(
            (a for a in anchors if a["text"].startswith("ach m1") and month_abbr in a["text"]),
            None,
        )
        ach_m1_total_idx, ach_m1_mm_idx, ach_m1_tr_idx = (
            scan_block(ach_m1_anchor["idx"], "total m1") if ach_m1_anchor else (-1, -1, -1)
        )

        targets = {}
        trustworthy = True
        for r in range(g_idx + 2, len(rows2d)):
            row = rows2d[r]
            if not row:
                continue
            is_total_row = any(v is not None and str(v).strip().lower() == "total" for v in row)
            if is_total_row:
                if all_rtr_idx == -1 or not (num(get(row, all_rtr_idx)) > 0):
                    trustworthy = False
                break
            cm_name = str(get(row, cm_idx) or "").strip()
            if not cm_name:
                continue
            val = num(get(row, total_ga_idx))
            if val > 0:
                targets[cm_name] = {
                    "total": val,
                    "massMigrant": num(get(row, mass_migrant_idx)) if mass_migrant_idx != -1 else 0,
                    "tourist": num(get(row, tourist_idx)) if tourist_idx != -1 else 0,
                    "m1Total": num(get(row, total_m1_idx)) if total_m1_idx != -1 else 0,
                    "m1MassMigrant": num(get(row, mm_m1_idx)) if mm_m1_idx != -1 else 0,
                    "m1Tourist": num(get(row, tr_m1_idx)) if tr_m1_idx != -1 else 0,
                    "achM1Total": num(get(row, ach_m1_total_idx)) if ach_m1_total_idx != -1 else 0,
                    "achM1MassMigrant": num(get(row, ach_m1_mm_idx)) if ach_m1_mm_idx != -1 else 0,
                    "achM1Tourist": num(get(row, ach_m1_tr_idx)) if ach_m1_tr_idx != -1 else 0,
                }

        if trustworthy and targets:
            results[month_abbr] = targets

    return results or None


def parse_mapping_sheets(map_sheets):
    """Builds codeToCanonical (merges Dtac+True codes of the same retailer),
    canonicalToCM, canonicalToMallType, canonicalToStock, targetByMonth,
    canonicalToName, canonicalToDistrict from every 'RTR Target GA...'
    mapping sheet found in the workbook. Ported from the mapCandidates
    .forEach loop in processRtrWorkbook, including the duplicate-month-block
    dedup (see BUGFIX comment inline). canonicalToName/District are the
    Python port's own addition (not in the original JS) -- needed so a
    roster store with zero sales in a given month (see
    fill_dormant_roster_stores) still has a name/district to publish, since
    that month's own sheet has no row for it to pull one from.

    Stock (T Stock/D Stock/Stock/Cover Mth) is a live snapshot at publish
    time, not month-keyed like target -- these columns appear once per
    mapping sheet, not once per month block, so no anchor-scanning needed,
    just a plain header lookup same as cluster/cm/mall type."""
    code_to_canonical = {}
    canonical_to_cm = {}
    canonical_to_mall_type = {}
    canonical_to_stock = {}  # {canonicalCode: {tStock, dStock, stock, coverMth}}
    canonical_to_name = {}  # {canonicalCode: retailer name} -- for roster stores with no
    # sales in a given month (see fill_dormant_roster_stores), which otherwise have no row
    # anywhere in that month's sheet to pull a name from.
    canonical_to_district = {}
    target_by_month = {}  # {monthAbbr: {canonicalCode: {target, m1}}}
    # raw code (Dtac or True) -> every canonical it's ever been paired to, in processing
    # order -- lets the Dtac/True merge below notice when a store's partner code changed
    # between mapping sheets (see the retirement block below).
    raw_seen_canonicals = {}

    for rows2d in map_sheets:
        h_idx = find_header_row(rows2d, MAP_REQUIRED_NORM, 20)
        if h_idx < 0:
            continue
        header = rows2d[h_idx]
        idx = {norm_header(h): i for i, h in enumerate(header)}
        cluster_col = idx.get("cluster")
        cm_col = idx.get("cm")
        mall_col = idx.get("in mall / out mall")
        name_col = idx.get("retailer name")
        district_col = idx.get("อำเภอ")
        t_stock_col = idx.get("t stock")
        d_stock_col = idx.get("d stock")
        stock_col = idx.get("stock")
        cover_col = idx.get("cover mth")
        blocks = parse_target_blocks(header)
        block_maps = [{} for _ in blocks]
        block_score = [0] * len(blocks)

        for r in range(h_idx + 1, len(rows2d)):
            row = rows2d[r]
            if not row:
                continue
            cluster = get(row, cluster_col) if cluster_col is not None else ""
            if not in_target_cluster(cluster):
                continue
            code_d = str(get(row, 0)).strip() if get(row, 0) not in (None, "") else None
            code_t = str(get(row, 1)).strip() if get(row, 1) not in (None, "") else None
            canonical = code_d or code_t
            if not canonical:
                continue

            # A store's Dtac/True partner code can differ between mapping sheets (e.g. its
            # pairing changed between the snapshots each sheet represents) -- when a raw
            # code we've already seen paired with a DIFFERENT canonical shows up again
            # paired with this one, the earlier canonical is a stale echo of this same
            # physical store, not a second store. Retire it (and redirect its own
            # self-reference) rather than leave it double-counted -- detected purely from
            # the pairing history here, not any fixed list, so this keeps working
            # whichever specific codes happen to conflict on a future refresh.
            for raw in (code_d, code_t):
                if not raw:
                    continue
                for prior_canonical in raw_seen_canonicals.get(raw, ()):
                    if prior_canonical != canonical and prior_canonical in canonical_to_cm:
                        del canonical_to_cm[prior_canonical]
                        canonical_to_mall_type.pop(prior_canonical, None)
                        canonical_to_stock.pop(prior_canonical, None)
                        canonical_to_name.pop(prior_canonical, None)
                        canonical_to_district.pop(prior_canonical, None)
                        code_to_canonical[prior_canonical] = canonical
                raw_seen_canonicals.setdefault(raw, set()).add(canonical)

            if code_d:
                code_to_canonical[code_d] = canonical
            if code_t:
                code_to_canonical[code_t] = canonical

            cm = str(get(row, cm_col)).strip() if cm_col is not None and get(row, cm_col) else UNMAPPED_CM
            canonical_to_cm[canonical] = cm
            mall_type = str(get(row, mall_col)).strip() if mall_col is not None and get(row, mall_col) else ""
            if mall_type:
                canonical_to_mall_type[canonical] = mall_type
            name = str(get(row, name_col)).strip() if name_col is not None and get(row, name_col) else ""
            if name:
                canonical_to_name[canonical] = name
            district = str(get(row, district_col)).strip() if district_col is not None and get(row, district_col) else ""
            if district:
                canonical_to_district[canonical] = DISTRICT_TH_TO_EN.get(district, district)

            if stock_col is not None or t_stock_col is not None or d_stock_col is not None:
                canonical_to_stock[canonical] = {
                    "tStock": num(get(row, t_stock_col)) if t_stock_col is not None else 0,
                    "dStock": num(get(row, d_stock_col)) if d_stock_col is not None else 0,
                    "stock": num(get(row, stock_col)) if stock_col is not None else 0,
                    "coverMth": num(get(row, cover_col)) if cover_col is not None else 0,
                }

            for bi, b in enumerate(blocks):
                rec = {
                    "target": num(get(row, b["targetIdx"])) if b["targetIdx"] >= 0 else 0,
                    "m1": num(get(row, b["m1Idx"])) if b["m1Idx"] >= 0 else 0,
                }
                block_maps[bi][canonical] = rec
                # Scored on EITHER target or m1 being set, not target alone -- target and m1
                # are entered independently, on their own schedule, within the same month's
                # block (confirmed against real data: September's target column was 0 for
                # every row while its m1 column had real values for 90/115 stores). Scoring
                # on target alone made a month with real m1 but not-yet-published targets
                # look "empty" by this heuristic's own definition, discarding the real m1
                # data right along with the genuinely-missing target.
                if rec["target"] > 0 or rec["m1"] > 0:
                    block_score[bi] += 1

        # BUGFIX (ported as-is): these sheets repeat the SAME month more than once, sometimes
        # as a near-empty leftover block. Pick the best-populated block per month rather than
        # trusting position, so a re-ordered or extra appendix block can't hijack the target.
        best_for_month = {}
        for bi, b in enumerate(blocks):
            cur = best_for_month.get(b["month"])
            if cur is None or block_score[bi] > block_score[cur]:
                best_for_month[b["month"]] = bi
        for mo, bi in best_for_month.items():
            if not block_score[bi]:
                continue  # block has neither target nor m1 for anyone -- genuinely empty
            target_by_month.setdefault(mo, {})
            # merged across sheets: MM and Tourist cover different stores, both count
            target_by_month[mo].update(block_maps[bi])

    return (code_to_canonical, canonical_to_cm, canonical_to_mall_type, canonical_to_stock,
            target_by_month, canonical_to_name, canonical_to_district)


def extract_routing_overrides(path, code_to_canonical):
    """Reads the 'routing' workbook's own 'RTR_Route' sheet for two things the primary
    workbook doesn't have: a genuinely live per-store target for whichever month this file is
    currently tracking, and each store's COM (field rep one level below CM -- confirmed by
    checking actual cardinality: the sheet's 'RS Name'/'PBH Name' columns are constant across
    the whole territory, but 'COM' varies, 4 distinct people each running a real subset of
    stores under one of the 2 CMs).

    The target month is found the same dynamic way parse_target_blocks already finds a 'GA by
    RSR <Month> Target' anchor in the primary workbook's own mapping sheets -- no month name
    is ever hardcoded, so this keeps working whichever month the file currently carries.

    Also reads each store's CM from this same sheet (the column after COM -- two columns in
    this sheet's header both normalize to "cm"; dict constructor below keeps the later/finer
    one as idx["cm"]). The primary workbook's own roster-mapping sheets are a point-in-time
    snapshot too (just an older one), so when the org chart changes -- a CM's book splitting
    into two new CMs, say -- the routing workbook catches up first. Call sites apply this CM
    override only to the month(s) this routing workbook is already the data source for (see
    process_rtr_workbook), never retroactively to older months the primary workbook still
    correctly covers, since this file has no way to say what was true in an earlier month.

    Returns (target_overrides, canonical_to_com, canonical_to_cm):
      target_overrides: {monthAbbr: {canonicalCode: target}}
      canonical_to_com: {canonicalCode: comName}
      canonical_to_cm: {canonicalCode: cmName}
    All {} if the sheet isn't shaped as expected -- this is optional extra data, a publish
    must still succeed without it (e.g. the file briefly missing between quarters)."""
    target_overrides = {}
    canonical_to_com = {}
    canonical_to_cm = {}
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        route_names = [n for n in wb.sheetnames if ROUTING_ROUTE_SHEET_RE.match(n.strip())]
        if not route_names:
            wb.close()
            return target_overrides, canonical_to_com, canonical_to_cm
        rows2d = sheet_to_rows(wb[route_names[0]])
        wb.close()
    except Exception as e:
        print(f"  (could not read routing workbook, non-fatal: {e})")
        return target_overrides, canonical_to_com, canonical_to_cm

    h_idx = find_header_row(rows2d, MAP_REQUIRED_NORM, 20)
    if h_idx < 0:
        return target_overrides, canonical_to_com, canonical_to_cm
    header = rows2d[h_idx]
    idx = {norm_header(h): i for i, h in enumerate(header)}
    cluster_col = idx.get("cluster")
    com_col = idx.get("com")
    cm_col = idx.get("cm")
    blocks = parse_target_blocks(header)

    for r in range(h_idx + 1, len(rows2d)):
        row = rows2d[r]
        if not row:
            continue
        cluster = get(row, cluster_col) if cluster_col is not None else None
        if not in_target_cluster(cluster):
            continue
        code_d = str(get(row, 0)).strip() if get(row, 0) not in (None, "") else None
        code_t = str(get(row, 1)).strip() if get(row, 1) not in (None, "") else None
        canonical = (code_to_canonical.get(code_d) or code_to_canonical.get(code_t)) if (code_d or code_t) else None
        if not canonical:
            continue  # not a code the primary workbook's own mapping already knows -- skip

        if com_col is not None:
            com = str(get(row, com_col) or "").strip()
            if com:
                canonical_to_com[canonical] = com
        if cm_col is not None:
            cm = str(get(row, cm_col) or "").strip()
            if cm:
                canonical_to_cm[canonical] = cm

        for b in blocks:
            tgt = num(get(row, b["targetIdx"])) if b["targetIdx"] >= 0 else 0
            if tgt > 0:
                target_overrides.setdefault(b["month"], {})[canonical] = tgt

    return target_overrides, canonical_to_com, canonical_to_cm


def new_daily_bucket():
    return {"ga": 0, "revAmt": 0, "mass": 0, "migrant": 0, "tourist": 0}


def aggregate_rtr_sheet(rows2d, code_to_canonical, canonical_to_cm, canonical_to_mall_type, canonical_to_stock=None,
                         canonical_to_name=None):
    """Per-store accumulation from one monthly RTR sheet. Ported from
    aggregateRtrSheet -- cluster filter, 7-11-named-row routing into a side
    channel (not RTR totals), Dtac/True code merging, tier columns, CM
    rollup. Returns None (skipped) if required columns are missing."""
    if not rows2d or len(rows2d) < 2:
        return None
    header = rows2d[0]
    idx = {h: i for i, h in enumerate(header) if h is not None}
    if not all(c in idx for c in REQUIRED_COLS):
        return None

    has_tier = [tc in idx for tc in TIER_COLS]
    stores = {}
    totals = {"ga": 0, "ap1d": 0, "ap1dAmt": 0, "ap30d": 0, "revAmt": 0, "mass": 0, "migrant": 0, "tourist": 0}
    daily = {}
    cm_rollup = {}
    row_count = 0
    unmapped_codes = set()
    unmapped_ga = 0

    for r in range(1, len(rows2d)):
        row = rows2d[r]
        if not row:
            continue
        cluster = get(row, idx["D_CLUSTER"])
        if not in_target_cluster(cluster):
            continue
        code = get(row, idx["PARTNER_CODE"])
        if code is None or code == "":
            continue
        code_str = str(code).strip()
        ga = num(get(row, idx["GA"]))
        ap1d = num(get(row, idx.get("AP1D", -1)))
        ap1d_amt = num(get(row, idx.get("AP1D_AMT", -1)))
        ap30d = num(get(row, idx.get("AP30D", -1)))
        rev_amt = num(get(row, idx.get("REV_AMT", -1)))
        seg = str(get(row, idx["GROUP_SIM"]) or "").upper()
        # str(...) is load-bearing, not cosmetic -- openpyxl reads TM_KEY_DAY as a Python
        # int when the cell is number-formatted and as a str when it's text-formatted, and
        # real RTR monthly sheets mix both within the same column. Without this coercion,
        # the exact same calendar day (e.g. int 20260401 vs str "20260401") becomes TWO
        # separate dict keys below, silently splitting that day's GA between them instead
        # of accumulating into one entry. Matches the SEV side's own TM_KEY_DAY handling,
        # which already does this (see aggregate_sev_rows).
        day = str(get(row, idx["TM_KEY_DAY"]))
        pname = get(row, idx["PARTNER_NAME"])

        # 7-11-named stores are a different channel -- route their transactions away from
        # RTR totals/CM mapping entirely. The publish pipeline doesn't need this side
        # channel's *contents* (the 7-11 tab is built from its own separate xlsx files),
        # only that these rows must NOT be double-counted into RTR -- so just skip them.
        if pname and SEV_NAME_RE.search(str(pname)):
            continue

        canonical_code = code_to_canonical.get(code_str, code_str)
        cm = canonical_to_cm.get(canonical_code, UNMAPPED_CM)
        if cm == UNMAPPED_CM:
            unmapped_codes.add(canonical_code)
            unmapped_ga += ga
            continue
        row_count += 1

        totals["ga"] += ga
        totals["ap1d"] += ap1d
        totals["ap1dAmt"] += ap1d_amt
        totals["ap30d"] += ap30d
        totals["revAmt"] += rev_amt
        if seg == "MASS":
            totals["mass"] += ga
        elif seg == "MIGRANT":
            totals["migrant"] += ga
        elif seg == "TOURIST":
            totals["tourist"] += ga

        d = daily.setdefault(day, new_daily_bucket())
        d["ga"] += ga
        d["revAmt"] += rev_amt
        if seg == "MASS":
            d["mass"] += ga
        elif seg == "MIGRANT":
            d["migrant"] += ga
        elif seg == "TOURIST":
            d["tourist"] += ga

        cr = cm_rollup.setdefault(cm, {
            "ga": 0, "ap1d": 0, "ap1dAmt": 0, "ap30d": 0, "revAmt": 0,
            "mass": 0, "migrant": 0, "tourist": 0, "storeCodes": set(), "daily": {},
        })
        cr["ga"] += ga
        cr["ap1d"] += ap1d
        cr["ap1dAmt"] += ap1d_amt
        cr["ap30d"] += ap30d
        cr["revAmt"] += rev_amt
        if seg == "MASS":
            cr["mass"] += ga
        elif seg == "MIGRANT":
            cr["migrant"] += ga
        elif seg == "TOURIST":
            cr["tourist"] += ga
        cr["storeCodes"].add(canonical_code)
        cr_d = cr["daily"].setdefault(day, new_daily_bucket())
        cr_d["ga"] += ga
        cr_d["revAmt"] += rev_amt
        if seg == "MASS":
            cr_d["mass"] += ga
        elif seg == "MIGRANT":
            cr_d["migrant"] += ga
        elif seg == "TOURIST":
            cr_d["tourist"] += ga

        s = stores.get(canonical_code)
        if s is None:
            # hasStock distinguishes "genuinely tracked at zero" from "this store just isn't
            # in the stock-mapping sheet at all" -- same reasoning as hasTarget below. Without
            # it, every store missing from the stock columns would default to coverMth=0 and
            # get flagged as critically low on stock, when the honest answer is "no data".
            stock_rec = (canonical_to_stock or {}).get(canonical_code)
            # Retailer Name (mapping sheet) over PARTNER_NAME (monthly sheet) -- PARTNER_NAME
            # is sometimes the SALESPERSON's own name, not the shop's (confirmed against real
            # data: e.g. "คุณสิริกร เรืองฤทธิ์" in PARTNER_NAME for a store the mapping sheet
            # calls "Lucky"). Falls back to PARTNER_NAME only for a code the mapping sheet
            # has no Retailer Name for at all.
            mapped_name = (canonical_to_name or {}).get(canonical_code)
            s = {
                "code": canonical_code, "name": mapped_name or pname or "Unknown", "cm": cm,
                "district": get(row, idx.get("DISTRICT_EN", -1)) or "",
                "mallType": canonical_to_mall_type.get(canonical_code, ""),
                "ga": 0, "ap1d": 0, "ap1dAmt": 0, "ap30d": 0, "revAmt": 0,
                "mass": 0, "migrant": 0, "tourist": 0,
                "tiers": [0] * 8, "simTypes": {}, "daily": {},
                "target": 0, "m1": 0,
                "hasStock": stock_rec is not None,
                "tStock": stock_rec["tStock"] if stock_rec else 0,
                "dStock": stock_rec["dStock"] if stock_rec else 0,
                "stock": stock_rec["stock"] if stock_rec else 0,
                "coverMth": stock_rec["coverMth"] if stock_rec else 0,
            }
            stores[canonical_code] = s
        s_d = s["daily"].setdefault(day, new_daily_bucket())
        s_d["ga"] += ga
        s_d["revAmt"] += rev_amt
        if seg == "MASS":
            s_d["mass"] += ga
        elif seg == "MIGRANT":
            s_d["migrant"] += ga
        elif seg == "TOURIST":
            s_d["tourist"] += ga
        s["ga"] += ga
        s["ap1d"] += ap1d
        s["ap1dAmt"] += ap1d_amt
        s["ap30d"] += ap30d
        s["revAmt"] += rev_amt
        if seg == "MASS":
            s["mass"] += ga
        elif seg == "MIGRANT":
            s["migrant"] += ga
        elif seg == "TOURIST":
            s["tourist"] += ga
        for t, tc in enumerate(TIER_COLS):
            if has_tier[t]:
                s["tiers"][t] += num(get(row, idx[tc]))
        sim_type = get(row, idx.get("SIM_TYPE", -1))
        if sim_type:
            s["simTypes"][sim_type] = s["simTypes"].get(sim_type, 0) + ga

    for cr in cm_rollup.values():
        cr["storeCount"] = len(cr["storeCodes"])
        del cr["storeCodes"]

    return {
        "rowCount": row_count,
        "totals": totals,
        "daily": daily,
        "cmRollup": cm_rollup,
        "stores": list(stores.values()),
        "unmappedCount": len(unmapped_codes),
        "unmappedGA": unmapped_ga,
    }


def fill_dormant_roster_stores(result, canonical_to_cm, canonical_to_name, canonical_to_district,
                                canonical_to_mall_type, canonical_to_stock):
    """Adds a zero-valued store entry for every roster code (canonical_to_cm) that had no
    sales row at all in this month's sheet. Without this, a store's total absence from
    RTR's "stores" table is indistinguishable from it simply not existing, which breaks
    two things that read that table directly: the CM leaderboard's store count undercounts
    (only counts stores that happened to sell something that month, not the CM's actual
    roster), and the "ยังไม่มียอดขายเลย" insight card (which filters for ga === 0) can
    never find anything, since a store with truly zero sales never got a row to begin
    with. Matches aggregate_rtr_sheet's own store shape exactly, just all-zero, so
    build_rtr_month_tables/resolve_rtr_targets don't need to know these are synthetic --
    call this BEFORE resolve_rtr_targets so a dormant store with a real committed target
    still gets it attached like any other store."""
    existing_codes = {s["code"] for s in result["stores"]}
    for canonical_code, cm in canonical_to_cm.items():
        if canonical_code in existing_codes:
            continue
        stock_rec = canonical_to_stock.get(canonical_code)
        result["stores"].append({
            "code": canonical_code,
            "name": canonical_to_name.get(canonical_code) or "Unknown",
            "cm": cm,
            "district": canonical_to_district.get(canonical_code, ""),
            "mallType": canonical_to_mall_type.get(canonical_code, ""),
            "ga": 0, "ap1d": 0, "ap1dAmt": 0, "ap30d": 0, "revAmt": 0,
            "mass": 0, "migrant": 0, "tourist": 0,
            "tiers": [0] * 8, "simTypes": {}, "daily": {},
            "target": 0, "m1": 0,
            "hasStock": stock_rec is not None,
            "tStock": stock_rec["tStock"] if stock_rec else 0,
            "dStock": stock_rec["dStock"] if stock_rec else 0,
            "stock": stock_rec["stock"] if stock_rec else 0,
            "coverMth": stock_rec["coverMth"] if stock_rec else 0,
        })


def resolve_rtr_targets(months, target_by_month, kpi_setup_targets):
    """Target-only subset of finalizeRtrMonths -- per-store target/hasTarget/m1
    from the mapping sheets, then a per-CM KPI Set Up override where one
    exists for that month, then month totals summed from the (post-override)
    CM rollups. Deliberately does NOT port achievePct/momPct/GA(LMTD) -- the
    dashboard only ever recomputes those client-side for display, from the
    raw numbers this script writes, so porting that math here would be pure
    unused risk. See finalizeRtrMonths in the dashboard JS for the full
    original (display-inclusive) version this is trimmed from."""
    for m in months:
        d = parse_sheet_date(m["sheetName"])
        mo_key = MONTH_LABEL[d["monthIndex"]].lower() if d else None
        tgt_map = target_by_month.get(mo_key, {}) if mo_key else {}

        for s in m["stores"]:
            t = tgt_map.get(s["code"])
            s["target"] = t["target"] if t else 0
            s["m1"] = t["m1"] if t else 0
            s["hasTarget"] = s["target"] > 0

        for cr in m["cmRollup"].values():
            cr["target"] = 0
            cr["m1"] = 0
            cr["targetedGA"] = 0
            cr["targetOverridden"] = False
            # 0, not None -- a CM with no KPI Set Up override this month simply has no
            # segment-scoped target either, same "no data" convention as target itself.
            cr["targetMassMigrant"] = 0
            cr["targetTourist"] = 0
            # Committed M1 target -- NOT the same thing as cr["m1"] (actual M1 achieved,
            # summed from stores just below). Only ever comes from the KPI Set Up override
            # (no per-store M1 target exists anywhere), same "0 = no data" convention.
            cr["targetM1"] = 0
            # Raw "Ach M1 (<Month>)" from the KPI Set Up sheet -- kept SEPARATE from cr["m1"]
            # (which stays the per-store sum, below) specifically so build_rtr_month_tables can
            # publish this raw figure as its own column. The published Sheet never carries
            # cr["m1"] itself -- the dashboard always re-sums it client-side from each store's
            # own m1 column -- so the override has to travel as its own signal the client can
            # apply the identical way, not as an already-resolved number Python keeps locally.
            cr["achM1"] = 0
        for s in m["stores"]:
            cr = m["cmRollup"].get(s["cm"])
            if cr is None:
                continue
            cr["m1"] += s["m1"]
            if s["hasTarget"]:
                cr["target"] += s["target"]
                cr["targetedGA"] += s["ga"]

        override_map = (kpi_setup_targets or {}).get(mo_key) if mo_key else None
        if override_map:
            for cm_name, cr in m["cmRollup"].items():
                if cm_name in override_map:
                    ov = override_map[cm_name]
                    cr["target"] = ov["total"]
                    cr["targetMassMigrant"] = ov["massMigrant"]
                    cr["targetTourist"] = ov["tourist"]
                    cr["targetM1"] = ov.get("m1Total", 0)
                    cr["targetOverridden"] = True
                    cr["achM1"] = ov.get("achM1Total", 0)

        # Sum from cmRollup (post-override), NOT independently from m["stores"] -- the KPI Set
        # Up override above replaces a CM's target/targetedGA with its committed, whole-
        # territory figures, and totals must reflect that too (the exact bug found and fixed
        # in the dashboard's own finalizeRtrMonths -- ported here the same corrected way).
        m["totals"]["target"] = sum(cr["target"] for cr in m["cmRollup"].values())
        m["totals"]["targetM1"] = sum(cr["targetM1"] for cr in m["cmRollup"].values())
        # "Ach M1" (by request) in place of the per-store sum wherever this sheet has a real
        # figure for that CM -- the dashboard applies this exact same resolution client-side
        # from the published cmtargets table, so Python's own totals here must match it.
        m["totals"]["m1"] = sum(
            cr["achM1"] if cr["achM1"] > 0 else cr["m1"] for cr in m["cmRollup"].values()
        )


def build_rtr_month_tables(m):
    """Output contract -- exact column order ported from buildRtrMonthTables.
    A cloud reader (readRtrMonthGroup) parses these back by fixed column
    index, so the order here must not change without changing that too."""
    stores_rows = []
    daily_rows = []
    sim_rows = []
    for s in m["stores"]:
        stores_rows.append([
            s["code"], s["name"], s["cm"], s["district"], s["mallType"],
            s["ga"], s["ap1d"], s["ap1dAmt"], s["ap30d"], s["revAmt"],
            s["mass"], s["migrant"], s["tourist"], s["target"], s["m1"],
        ] + s["tiers"] + [
            # Appended after tiers, not inserted earlier -- readRtrMonthGroup (the cloud reader)
            # parses this row by FIXED COLUMN INDEX, so existing columns must never move.
            # 1/0 not True/False: this becomes a Sheets cell, and Python bool would write the
            # literal words "True"/"False" there instead of a usable 0/1.
            1 if s["hasStock"] else 0,
            s["tStock"], s["dStock"], s["stock"], s["coverMth"],
            s.get("com", ""),
        ])
        for day in sorted(s["daily"].keys(), key=str):
            d = s["daily"][day]
            daily_rows.append([s["code"], day, d["ga"], d["revAmt"], d["mass"], d["migrant"], d["tourist"]])
        for st, v in s["simTypes"].items():
            sim_rows.append([s["code"], st, v])

    cm_rows = sorted(
        [
            [cm, cr["target"], cr["targetMassMigrant"], cr["targetTourist"], cr["targetM1"],
             cr["achM1"]]
            for cm, cr in m["cmRollup"].items() if cr["targetOverridden"]
        ],
        key=lambda r: r[0],
    )

    return {
        "meta": {
            "headers": ["field", "value"],
            "rows": [["sheetName", m["sheetName"]], ["label", m["label"]], ["rowCount", m["rowCount"]]],
        },
        "stores": {
            "headers": ["code", "name", "cm", "district", "mallType", "ga", "ap1d", "ap1dAmt",
                        "ap30d", "revAmt", "mass", "migrant", "tourist", "target", "m1"]
                       + ["tier_" + l for l in TIER_LABELS]
                       + ["hasStock", "tStock", "dStock", "stock", "coverMth", "com"],
            "rows": stores_rows,
        },
        "daily": {
            "headers": ["code", "date", "ga", "revAmt", "mass", "migrant", "tourist"],
            "rows": daily_rows,
        },
        "simtypes": {"headers": ["code", "simType", "ga"], "rows": sim_rows},
        "cmtargets": {
            "headers": ["cm", "target", "targetMassMigrant", "targetTourist", "targetM1", "achM1"],
            "rows": cm_rows,
        },
        # Per-COM committed target/M1 from the routing workbook's "Performance COM" sheet --
        # empty for any month that sheet doesn't cover (the dashboard then falls back to
        # summing per-store targets, as before).
        "comtargets": {
            "headers": ["com", "target", "targetMassMigrant", "targetTourist", "targetM1", "achM1"],
            "rows": sorted(
                [com, v["total"], v["massMigrant"], v["tourist"], v["m1Total"], v["achM1Total"]]
                for com, v in (m.get("comTargets") or {}).items()
            ),
        },
    }


def process_rtr_workbook(path, sheet_overrides=None, routing_path=None):
    # sheet_overrides: optional {month_sheet_name: rows2d} to substitute in place of that sheet's
    # own rows (read from a DIFFERENT workbook via sheet_to_rows there) -- e.g. a fresher pull of
    # the same month from a newer file that doesn't carry the mapping/KPI Set Up sheets needed to
    # process a whole workbook on its own. Every other month, and all mapping/target resolution,
    # still comes from `path` unchanged. None/{} (the default) is the normal, unmodified path.
    #
    # routing_path: optional path to the "routing" workbook (see _find_routing_workbook/
    # extract_routing_overrides) -- when given, its freshest raw month sheet (by row count, not
    # a hardcoded name) is folded into sheet_overrides automatically, its per-store target for
    # whichever month it's tracking patches in AFTER resolve_rtr_targets (the primary workbook's
    # own target for that same month can genuinely be all-zero -- not a bug, see the function's
    # own docstring), and every store gets its COM (field rep one level below CM) attached.
    sheet_overrides = dict(sheet_overrides or {})
    print(f"Reading RTR workbook: {path.name}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    all_names = wb.sheetnames
    month_candidates = [n for n in all_names if MONTH_RE.match(n.strip())]
    map_candidates = [n for n in all_names if MAP_SHEET_RE.search(n)]
    kpi_candidates = [n for n in all_names if KPI_SETUP_RE.search(n)]
    print(f"  {len(all_names)} sheets total: {len(month_candidates)} monthly, "
          f"{len(map_candidates)} mapping, {len(kpi_candidates)} KPI Set Up")
    if not month_candidates:
        raise SystemExit(f"No monthly sheets found in {path.name} (expected names like 'Aug.26')")

    kpi_setup_targets = None
    for name in kpi_candidates:
        rows2d = sheet_to_rows(wb[name])
        kpi_setup_targets = parse_kpi_setup_targets(rows2d)
        if kpi_setup_targets:
            break

    map_sheets = [sheet_to_rows(wb[name]) for name in map_candidates]
    (code_to_canonical, canonical_to_cm, canonical_to_mall_type, canonical_to_stock,
     target_by_month, canonical_to_name, canonical_to_district) = parse_mapping_sheets(map_sheets)
    print(f"  mapping: {len(code_to_canonical)} codes -> {len(canonical_to_cm)} roster stores, "
          f"{len(set(canonical_to_cm.values()))} CMs, "
          f"{len(canonical_to_stock)} stores with stock data, "
          f"targets resolved for months: {sorted(target_by_month.keys())}")
    if kpi_setup_targets:
        print(f"  KPI Set Up override available for months: {sorted(kpi_setup_targets.keys())}")

    routing_target_overrides = {}
    canonical_to_com = {}
    canonical_to_cm_override = {}
    com_targets = {}  # {monthAbbr: {comName: {total, massMigrant, tourist, m1Total, achM1Total, ...}}}
    # Month sheet NAMES the routing workbook's CM override is trusted for -- deliberately
    # NARROWER than "months whose raw rows came from the routing workbook". A month the
    # routing workbook merely has a row-count-fresher copy of (its own "Sep" outscoring the
    # primary workbook's "Sep.26") still had ITS org chart correctly described by the
    # primary workbook's own roster -- the routing file being a better source of that
    # month's SALES total doesn't make its (now-updated-for-the-newest-month) org chart
    # correct for that same month retroactively. Only a month the primary workbook doesn't
    # carry AT ALL can unambiguously only mean "right now", so routing's current org chart
    # is only trusted there.
    routing_cm_months = set()
    if routing_path:
        print(f"  routing workbook: {routing_path.name}")
        routing_target_overrides, canonical_to_com, canonical_to_cm_override = extract_routing_overrides(
            routing_path, code_to_canonical)
        if canonical_to_com:
            print(f"    {len(canonical_to_com)} stores have a COM assignment")
        if routing_target_overrides:
            print(f"    per-store target override available for: {sorted(routing_target_overrides.keys())}")
        try:
            wb_r = openpyxl.load_workbook(routing_path, read_only=True, data_only=True)
            kpi_setup_targets = kpi_setup_targets or {}
            for pname in [n for n in wb_r.sheetnames if PERF_CM_SHEET_RE.match(n.strip())]:
                perf_targets = parse_kpi_setup_targets(sheet_to_rows(wb_r[pname])) or {}
                for mo, cms in perf_targets.items():
                    if mo not in kpi_setup_targets:
                        kpi_setup_targets[mo] = cms
                        print(f"    CM targets for {mo} from {pname} (primary KPI Set Up has no {mo} block)")
            for pname in [n for n in wb_r.sheetnames if PERF_COM_SHEET_RE.match(n.strip())]:
                for mo, coms in (parse_kpi_setup_targets(sheet_to_rows(wb_r[pname]), name_col="com_name")
                                 or {}).items():
                    com_targets.setdefault(mo, {}).update(coms)
                    print(f"    COM targets for {mo} from {pname}: {len(coms)} COMs")
            for rname in [n for n in wb_r.sheetnames if MONTH_RE.match(n.strip())]:
                r_date = parse_sheet_date(rname)
                if not r_date:
                    continue
                match = next((pn for pn in month_candidates
                              if (parse_sheet_date(pn) or {}).get("monthIndex") == r_date["monthIndex"]), None)
                r_rows = sheet_to_rows(wb_r[rname])
                if not match:
                    # A month the primary workbook doesn't carry yet at all (e.g. the routing
                    # file already has October while the primary still stops at September) --
                    # add it as its own month rather than silently dropping it.
                    print(f"    {rname} ({len(r_rows)} rows) is a month the primary workbook "
                          f"doesn't have yet -- adding it")
                    month_candidates.append(rname)
                    sheet_overrides[rname] = r_rows
                    routing_cm_months.add(rname)
                    continue
                if len(r_rows) > wb[match].max_row:
                    print(f"    {rname} ({len(r_rows)} rows) is fresher than {match} "
                          f"({wb[match].max_row} rows) -- using it instead")
                    sheet_overrides[match] = r_rows
            wb_r.close()
        except Exception as e:
            print(f"  (could not read routing workbook's monthly sheets, non-fatal: {e})")

    canonical_to_cm_routing = dict(canonical_to_cm)
    canonical_to_cm_routing.update(canonical_to_cm_override)
    if routing_cm_months and canonical_to_cm_override:
        print(f"    {len(canonical_to_cm_override)} stores' CM resolved from the routing workbook "
              f"for {sorted(routing_cm_months)} (primary roster still used for every other month)")

    months = []
    for name in month_candidates:
        # Each sheet is thousands of rows read via openpyxl read_only -- the slow part of this
        # whole script -- so this prints AS EACH ONE finishes rather than staying silent until
        # every sheet is done and dumping the results all at once.
        print(f"  reading {name}...", end="", flush=True)
        rows2d = sheet_overrides[name] if name in sheet_overrides else sheet_to_rows(wb[name])
        cm_map = canonical_to_cm_routing if name in routing_cm_months else canonical_to_cm
        result = aggregate_rtr_sheet(rows2d, code_to_canonical, cm_map, canonical_to_mall_type, canonical_to_stock, canonical_to_name)
        if result is None:
            print(f" skipped (missing required columns: {', '.join(REQUIRED_COLS)})")
            continue
        active_count = len(result["stores"])
        fill_dormant_roster_stores(result, cm_map, canonical_to_name, canonical_to_district,
                                    canonical_to_mall_type, canonical_to_stock)
        print(f" done ({result['rowCount']} rows, {active_count} active + "
              f"{len(result['stores']) - active_count} dormant = {len(result['stores'])} roster stores)")
        months.append({"sheetName": name, "label": parse_sheet_label(name), **result})
    wb.close()

    if not months:
        raise SystemExit(f"No usable monthly sheets in {path.name}")
    sort_months_chronologically(months)
    resolve_rtr_targets(months, target_by_month, kpi_setup_targets)

    for m in months:
        for s in m["stores"]:
            s["com"] = canonical_to_com.get(s["code"], "")
        d = parse_sheet_date(m["sheetName"])
        mo_key = MONTH_LABEL[d["monthIndex"]].lower() if d else None
        coms_here = {s["com"] for s in m["stores"] if s["com"]}
        m["comTargets"] = {com: v for com, v in com_targets.get(mo_key, {}).items() if com in coms_here}
    if routing_target_overrides:
        for m in months:
            d = parse_sheet_date(m["sheetName"])
            mo_key = MONTH_LABEL[d["monthIndex"]].lower() if d else None
            overrides = routing_target_overrides.get(mo_key) if mo_key else None
            if not overrides:
                continue
            patched = 0
            for s in m["stores"]:
                new_tgt = overrides.get(s["code"])
                if new_tgt:
                    s["target"] = new_tgt
                    s["hasTarget"] = True
                    patched += 1
            if patched:
                print(f"  {m['label']}: patched per-store target onto {patched} stores "
                      f"from the routing workbook")

    for m in months:
        t = m["totals"]
        print(f"  {m['label']:<10} GA {t['ga']:>8.0f}  Target {t['target']:>8.0f}  "
              f"Achieve {(t['ga']/t['target']*100) if t['target'] else 0:>5.1f}%  "
              f"({m['rowCount']} rows, {m['unmappedCount']} unmapped codes / GA {m['unmappedGA']:.0f})")
        for cm, cr in sorted(m["cmRollup"].items()):
            flag = " [KPI Set Up override]" if cr["targetOverridden"] else ""
            print(f"      {cm:<20} GA {cr['ga']:>8.0f}  Target {cr['target']:>8.0f}{flag}")

    return months


# ═════════════════════════════════════════════════════════════════════════
# 7-11 PIPELINE
# ═════════════════════════════════════════════════════════════════════════

def find_header_idx(idx, patterns):
    for p in patterns:
        if idx.get(p) is not None:
            return idx[p]
    # fallback: fuzzy-match any header containing all the pattern's normalized words
    # (ก-๙ = U+0E01-U+0E59, same Thai letter+digit range the JS pattern uses)
    for p in patterns:
        norm = re.sub(r"[^a-z0-9ก-๙]+", "", p.lower())
        for key, i in idx.items():
            key_norm = re.sub(r"[^a-z0-9ก-๙]+", "", key.lower())
            if key_norm == norm:
                return i
    return None


def extract_sim_price(sim_type):
    if not sim_type:
        return None
    s = str(sim_type)
    m = re.search(r"(\d+)\s*(?:B\b|THB)", s, re.I)
    if m:
        return int(m.group(1))
    if re.search(r"\dDAYS?\b", s, re.I):
        return None  # e.g. "8DAYS"/"15DAYS" -- a duration, not a price
    m2 = re.search(r"(\d+)\s*$", s)
    if m2:
        return int(m2.group(1))
    return None


def sim_exact_price_band(price, sim_type):
    if price == 49:
        return "p49"
    if price == 99:
        return "p99"
    if price == 199:
        return "p199"
    if price is not None and price > 199:
        return "pOver199"
    if price is None and sim_type and re.search(r"infinite", str(sim_type), re.I):
        return "pOver199"
    return None


def looks_like_roster_sheet(rows2d):
    header = rows2d[0] if rows2d else []
    trimmed = {str(h).strip() if h is not None else "" for h in header}
    return "CM Name" in trimmed and "TM_KEY_DAY" not in trimmed


def looks_like_sev_target_sheet(sheet_name):
    return bool(SEV_TARGET_SHEET_RE.match(str(sheet_name or "").strip()))


def looks_like_sev_cm_target_sheet(sheet_name):
    return bool(SEV_CM_TARGET_SHEET_RE.match(str(sheet_name or "").strip()))


def sev_count_rows_by_day(rows2d):
    counts = {}
    if not rows2d or len(rows2d) < 2:
        return counts
    header = rows2d[0]
    idx = {str(h).strip() if h is not None else "": i for i, h in enumerate(header)}
    day_idx, prov_idx = idx.get("TM_KEY_DAY"), idx.get("TDS_PROVINCE")
    if day_idx is None or prov_idx is None:
        return counts
    for r in range(1, len(rows2d)):
        row = rows2d[r]
        if not row or not in_target_cluster(get(row, prov_idx)):
            continue
        day = str(get(row, day_idx) if get(row, day_idx) is not None else "").strip()
        if len(day) < 6:
            continue
        counts[day] = counts.get(day, 0) + 1
    return counts


def sev_months_touched(ws, day_col_idx):
    """Cheap day-only scan of one sheet -- reads just the TM_KEY_DAY column (not the
    whole row) across every row, to find which months this sheet has ANY rows for at
    all, before committing to the full (much more expensive) sheet_to_rows +
    aggregate_sev_rows pass. Column-restricted iter_rows still streams row-by-row in
    openpyxl's read_only mode (unlike iter_cols, which read_only mode doesn't support),
    so this is real savings, not just deferred work. Doesn't apply the cluster/province
    filter aggregate_sev_rows does -- this only needs to answer "could this sheet matter
    at all for the months about to be kept", a broader question than "which of its
    cluster rows survive", so skipping that filter here can't cause anything to be
    wrongly excluded."""
    months = set()
    for row in ws.iter_rows(min_row=2, min_col=day_col_idx + 1, max_col=day_col_idx + 1, values_only=True):
        val = row[0] if row else None
        day = str(val).strip() if val is not None else ""
        # Require an actual 8-digit YYYYMMDD, not just "6+ characters" -- a pivot/summary
        # sheet can have a column literally named TM_KEY_DAY (inherited from a copied
        # pivot field) whose cells hold text like "BKK : ..." or "Row Labels" instead of
        # real dates. str(text)[:6] used to pass the old length check and get treated as a
        # bogus "month", which (being all-letters) sorts after every real YYYYMM string and
        # can hijack the newest-6-months window, starving out every sheet with real data.
        if len(day) == 8 and day.isdigit():
            months.add(day[:6])
    return months


def new_sev_store(code, name):
    return {
        "code": code, "name": name or "Unknown", "district": "", "ae": UNMAPPED_AE,
        "zone": "", "gm": "ไม่ระบุ GM", "fyai": "ไม่ระบุฝ่าย", "khet": "ไม่ระบุเขต",
        "fc": "", "tambon": "", "cm": None,
        "ga": 0, "revAmt": 0, "ap100": 0, "ap49": 0, "ap99": 0, "ap199": 0, "apOver199": 0,
        "mass": 0, "migrant": 0, "tourist": 0, "dtacGa": 0, "trueGa": 0,
        "daily": {}, "byMonth": {}, "simTypes": {},
        # Per month, not a single value -- a store can have a committed target from more
        # than one source at once (e.g. August's "TG Store" sheet AND September's "Monitor
        # Store" sheet, each covering its own month), and each needs to reach its own
        # month's published tab without one silently overwriting the other.
        "targetByMonth": {}, "aeMonthKey": None,
        # newest month s["cm"] was actually set from -- same reasoning/pattern as aeMonthKey:
        # starting with the new monthly "AE Monitoring" format, each month's own Data sheet
        # carries CM directly per row, and the newest month's assignment should win over an
        # older month's, not just whichever sheet happened to be processed last.
        "cmMonthKey": None,
    }


def fill_if_placeholder(current, placeholder, new_val):
    if new_val is None or new_val == "":
        return current
    trimmed = str(new_val).strip()
    return trimmed if current == placeholder and trimmed else current


def aggregate_sev_rows(rows2d, stores_map, stats, day_owner, sheet_idx):
    """Ported from aggregateSevRows -- per-row accumulation honoring
    dayOwner (so overlapping monthly files can't double-count a calendar
    day), cluster filter, AE/GM/ฝ่าย/เขต placeholder-backfill and
    latest-month-wins AE reassignment, SIM_TYPE-derived AP price bands."""
    if not rows2d or len(rows2d) < 2:
        return
    header = rows2d[0]
    idx = {str(h).strip() if h is not None else "": i for i, h in enumerate(header)}
    missing = [c for c in SEV_REQUIRED_HEADERS if c not in idx]
    if missing:
        raise ValueError("ขาดคอลัมน์: " + ", ".join(missing))

    rev_idx = find_header_idx(idx, ["Sum of REV_AMT", " Sum of REV_AMT "])
    ap100_idx = find_header_idx(idx, ["Sum of AP>100", "Sum of AP100UP", "Sum of AP 100UP"])
    ae_idx = find_header_idx(idx, ["AE Name", "COM"])
    cm_idx = find_header_idx(idx, ["CM Name", "CM"])

    for r in range(1, len(rows2d)):
        row = rows2d[r]
        if not row:
            continue
        province = get(row, idx["TDS_PROVINCE"])
        stats["totalRows"] += 1
        if not in_target_cluster(province):
            continue
        stats["clusterRows"] += 1

        code = get(row, idx["PARTNER_CODE"])
        if code is None or code == "":
            continue
        code_str = str(code).strip()
        if code_str in EXCLUDED_STORE_CODES:
            continue
        sim_type = get(row, idx.get("SIM_TYPE", -1))
        if sim_type and str(sim_type).strip().upper() in EXCLUDED_SIM_TYPES:
            continue
        day = str(get(row, idx["TM_KEY_DAY"])).strip()
        if day_owner is not None and day_owner.get(day) is not None and day_owner[day] != sheet_idx:
            continue  # a different sheet has more complete data for this exact day
        month_key = day[:6] if len(day) >= 6 else None
        ga = num(get(row, idx["Sum of GA"]))
        rev_amt = num(get(row, rev_idx)) if rev_idx is not None else 0
        ap100 = num(get(row, ap100_idx)) if ap100_idx is not None else 0
        seg = str(get(row, idx["GROUP_SIM"]) or "").upper()
        company = str(get(row, 0) or "").strip().upper()
        is_dtac = "DTAC" in company
        is_true = "TRUE" in company

        s = stores_map.get(code_str)
        if s is None:
            s = new_sev_store(code_str, get(row, idx["PARTNER_NAME"]))
            stores_map[code_str] = s

        if idx.get("DISTRICT_EN") is not None and not s["district"]:
            s["district"] = get(row, idx["DISTRICT_EN"]) or ""
        if ae_idx is not None and get(row, ae_idx):
            ae_val = str(get(row, ae_idx)).strip()
            if ae_val and (month_key is None or s["aeMonthKey"] is None or month_key >= s["aeMonthKey"]):
                s["ae"] = ae_val
                if month_key:
                    s["aeMonthKey"] = month_key
        if idx.get("Zone") is not None and not s["zone"]:
            s["zone"] = str(get(row, idx["Zone"]) or "").strip()
        s["gm"] = fill_if_placeholder(s["gm"], "ไม่ระบุ GM", get(row, idx.get("GM", -1)))
        s["fyai"] = fill_if_placeholder(s["fyai"], "ไม่ระบุฝ่าย", get(row, idx.get("ฝ่าย", -1)))
        s["khet"] = fill_if_placeholder(s["khet"], "ไม่ระบุเขต", get(row, idx.get("เขต", -1)))
        if idx.get("FC") is not None and not s["fc"]:
            s["fc"] = str(get(row, idx["FC"]) or "").strip()
        if idx.get("ตำบล") is not None and not s["tambon"]:
            s["tambon"] = str(get(row, idx["ตำบล"]) or "").strip()
        if cm_idx is not None and get(row, cm_idx):
            cm_val = str(get(row, cm_idx)).strip()
            if cm_val and (month_key is None or s["cmMonthKey"] is None or month_key >= s["cmMonthKey"]):
                s["cm"] = cm_val
                if month_key:
                    s["cmMonthKey"] = month_key

        s["ga"] += ga
        s["revAmt"] += rev_amt
        s["ap100"] += ap100
        if seg == "MASS":
            s["mass"] += ga
        elif seg == "MIGRANT":
            s["migrant"] += ga
        elif seg == "TOURIST":
            s["tourist"] += ga
        if is_dtac:
            s["dtacGa"] += ga
        elif is_true:
            s["trueGa"] += ga

        band = sim_exact_price_band(extract_sim_price(sim_type), sim_type)
        if band == "p49":
            s["ap49"] += ga
        elif band == "p99":
            s["ap99"] += ga
        elif band == "p199":
            s["ap199"] += ga
        elif band == "pOver199":
            s["apOver199"] += ga
        if sim_type:
            s["simTypes"][sim_type] = s["simTypes"].get(sim_type, 0) + ga

        d = s["daily"].setdefault(day, {"ga": 0, "revAmt": 0, "dtacGa": 0, "trueGa": 0,
                                         "ap49": 0, "ap99": 0, "ap199": 0, "apOver199": 0,
                                         "mass": 0, "migrant": 0, "tourist": 0})
        d["ga"] += ga
        d["revAmt"] += rev_amt
        if is_dtac:
            d["dtacGa"] += ga
        elif is_true:
            d["trueGa"] += ga
        if band == "p49":
            d["ap49"] += ga
        elif band == "p99":
            d["ap99"] += ga
        elif band == "p199":
            d["ap199"] += ga
        elif band == "pOver199":
            d["apOver199"] += ga
        if seg == "MASS":
            d["mass"] += ga
        elif seg == "MIGRANT":
            d["migrant"] += ga
        elif seg == "TOURIST":
            d["tourist"] += ga

        if month_key:
            bm = s["byMonth"].setdefault(month_key, {"ga": 0, "revAmt": 0, "ap100": 0,
                                                       "ap49": 0, "ap99": 0, "ap199": 0, "apOver199": 0,
                                                       "mass": 0, "migrant": 0, "tourist": 0,
                                                       "dtacGa": 0, "trueGa": 0, "simTypes": {}})
            bm["ga"] += ga
            bm["revAmt"] += rev_amt
            bm["ap100"] += ap100
            if band == "p49":
                bm["ap49"] += ga
            elif band == "p99":
                bm["ap99"] += ga
            elif band == "p199":
                bm["ap199"] += ga
            elif band == "pOver199":
                bm["apOver199"] += ga
            if seg == "MASS":
                bm["mass"] += ga
            elif seg == "MIGRANT":
                bm["migrant"] += ga
            elif seg == "TOURIST":
                bm["tourist"] += ga
            if is_dtac:
                bm["dtacGa"] += ga
            elif is_true:
                bm["trueGa"] += ga
            if sim_type:
                bm["simTypes"][sim_type] = bm["simTypes"].get(sim_type, 0) + ga


def aggregate_roster_rows(rows2d, code_to_cm, roster_stats):
    if not rows2d or len(rows2d) < 2:
        return
    header = rows2d[0]
    idx = {str(h).strip() if h is not None else "": i for i, h in enumerate(header)}
    if idx.get("PARTNER_CODE") is None or idx.get("CM Name") is None:
        raise ValueError("ไฟล์ roster ขาดคอลัมน์ PARTNER_CODE หรือ CM Name")
    for r in range(1, len(rows2d)):
        row = rows2d[r]
        if not row:
            continue
        code = get(row, idx["PARTNER_CODE"])
        if code is None or code == "":
            continue
        cm = get(row, idx["CM Name"])
        if cm is None or cm == "":
            continue
        code_to_cm[str(code).strip()] = str(cm).strip()
        roster_stats["rows"] += 1


def aggregate_sev_target_rows(rows2d, code_to_target, target_stats):
    h_idx = find_header_row(rows2d, SEV_TARGET_REQUIRED_NORM, 10)
    if h_idx < 0:
        return
    header = rows2d[h_idx]
    idx = {norm_header(h): i for i, h in enumerate(header)}
    code_idx = idx.get("partner_code")
    if code_idx is None:
        return

    group_row_idx = -1
    for gr in range(h_idx - 1, max(-1, h_idx - 6), -1):
        candidate = rows2d[gr] if gr < len(rows2d) else None
        if candidate and any(norm_header(v) == "overall" for v in candidate):
            group_row_idx = gr
            break
    if group_row_idx < 0:
        return
    group_row = rows2d[group_row_idx]
    overall_idx = next((g for g, v in enumerate(group_row) if norm_header(v) == "overall"), -1)
    if overall_idx < 0:
        return
    target_ga_idx = -1
    for c in range(overall_idx, len(header)):
        if norm_header(get(header, c)) == "gross add":
            target_ga_idx = c
            break
    if target_ga_idx < 0:
        return

    for r in range(h_idx + 1, len(rows2d)):
        row = rows2d[r]
        if not row:
            continue
        code = get(row, code_idx)
        if code is None or code == "":
            continue
        target = num(get(row, target_ga_idx))
        if target > 0:
            code_to_target[str(code).strip()] = target
            target_stats["rows"] += 1


def find_sev_monitor_store_group_row(rows2d, max_scan=10):
    """Finds the row carrying a Monitor Store-shaped sheet's Mass/Migrant/Tourist/All
    segment group labels (see aggregate_sev_monitor_store_targets), and the column each
    group's block starts at. Detected structurally, not by sheet name -- this shape has so
    far only shown up as AE Monitoring's own "Monitor Store" tab, but content-based
    detection means this keeps working even if that name changes or a similarly-shaped
    sheet turns up elsewhere, rather than being tied to a name that could drift."""
    for i in range(min(max_scan, len(rows2d))):
        row = rows2d[i]
        if not row:
            continue
        positions = {}
        for j, cell in enumerate(row):
            text = str(cell).strip() if cell is not None else ""
            if text in ("Mass", "Migrant", "Tourist", "All") and text not in positions:
                positions[text] = j
        if all(k in positions for k in ("Mass", "Migrant", "Tourist", "All")):
            return i, positions
    return -1, {}


def looks_like_sev_monitor_store_sheet(rows2d):
    group_row_idx, _ = find_sev_monitor_store_group_row(rows2d)
    return group_row_idx >= 0


def aggregate_sev_monitor_store_targets(rows2d, code_to_target, stats):
    """Monitor Store-shaped sheets (see find_sev_monitor_store_group_row) carry a per-store
    target for each of Mass/Migrant/Tourist, plus a combined "All" block that's their sum
    (verified against real data: All's Actual column exactly equals Mass+Migrant+Tourist's
    own Actual columns added together) -- the "All" block's target is the one used here,
    matching what a store's overall target means everywhere else in this pipeline. Each
    block is a fixed 10-column run starting at its group label's own column: +4 = target
    (mislabeled "TG Mass" in every block, not just the Mass one -- position, not the header
    text, is what actually distinguishes them), +5 = actual GA so far this month."""
    group_row_idx, positions = find_sev_monitor_store_group_row(rows2d)
    if group_row_idx < 0 or "All" not in positions:
        return
    header_row_idx = group_row_idx + 1
    if header_row_idx >= len(rows2d):
        return
    header = rows2d[header_row_idx]
    idx = {str(h).strip() if h is not None else "": i for i, h in enumerate(header)}
    code_idx = idx.get("PARTNER_CODE")
    if code_idx is None:
        return
    target_idx = positions["All"] + 4

    for r in range(header_row_idx + 1, len(rows2d)):
        row = rows2d[r]
        if not row:
            continue
        code = get(row, code_idx)
        if code is None or code == "":
            continue
        target = num(get(row, target_idx))
        if target > 0:
            code_to_target[str(code).strip()] = target
            stats["rows"] += 1


def find_sev_le_by_ae_columns(rows2d, max_scan=10):
    """Finds every "LE by AE" column in an "LE by store"-shaped sheet (see
    aggregate_sev_le_by_store_targets), keyed by the month abbreviation in the group-row
    label directly above each. This file repeats the "LE by AE" header once per month
    shown side by side (e.g. one column for Aug, one for Sep) -- the header text alone
    can't tell them apart, the same "position/label, not header text" pattern as Monitor
    Store's repeated "TG Mass" blocks. Returns (header_row_idx, {month_abbr: col_idx}), or
    (-1, {}) if this sheet isn't shaped like this at all."""
    for i in range(min(max_scan, len(rows2d)) - 1):
        group_row = rows2d[i]
        header_row = rows2d[i + 1]
        if not group_row or not header_row:
            continue
        cols = {}
        for j, h in enumerate(header_row):
            if h is None or norm_header(h) != "le by ae":
                continue
            month_label = str(group_row[j]).strip().lower() if j < len(group_row) and group_row[j] is not None else ""
            if month_label in MONTH_ABBR:
                cols[month_label] = j
        if cols:
            return i + 1, cols
    return -1, {}


def looks_like_sev_le_by_store_sheet(rows2d):
    header_row_idx, _ = find_sev_le_by_ae_columns(rows2d)
    return header_row_idx >= 0


def aggregate_sev_le_by_store_targets(rows2d, code_to_target, stats, month_abbr):
    """Per-store target for `month_abbr` from an "LE by store" file's own "LE by AE"
    column for that month -- see find_sev_le_by_ae_columns for the column-finding logic.
    "LE" (Latest Estimate) is this business's own name for a forecast/target figure,
    confirmed against real data: Pak Kret cluster stores show realistic per-store values
    (2-22 GA in the Sep'26 file), and stores marked Status=Close correctly show 0, not a
    stray leftover number."""
    header_row_idx, cols = find_sev_le_by_ae_columns(rows2d)
    if header_row_idx < 0 or month_abbr not in cols:
        return
    header = rows2d[header_row_idx]
    idx = {str(h).strip() if h is not None else "": i for i, h in enumerate(header)}
    code_idx = idx.get("PARTNER_CODE")
    if code_idx is None:
        return
    target_col = cols[month_abbr]

    for r in range(header_row_idx + 1, len(rows2d)):
        row = rows2d[r]
        if not row:
            continue
        code = get(row, code_idx)
        if code is None or code == "":
            continue
        target = num(get(row, target_col))
        if target > 0:
            code_to_target[str(code).strip()] = target
            stats["rows"] += 1


def find_sev_le_cm_target_block(rows2d, max_scan=10):
    """Finds the CM-level block in an "LE by store" file's "PV" sheet -- that sheet has
    TWO side-by-side pivot blocks sharing overlapping column names (Row Labels/CM Name/
    Sum of LE by AE/TG <month>/Diff): one at the AE level (has an "AE Name" column between
    CM Name and the Sum column) and one already rolled up to CM level (CM Name goes
    straight to the Sum/TG columns). Distinguished by that structural difference, not by
    which block comes first, since either could in principle be on the left."""
    for i in range(min(max_scan, len(rows2d))):
        row = rows2d[i]
        if not row:
            continue
        normed = [norm_header(v) for v in row]
        for j, h in enumerate(normed):
            if h != "cm name":
                continue
            next_h = normed[j + 1] if j + 1 < len(normed) else ""
            if next_h == "ae name":
                continue  # AE-level block, not the CM-level one we want
            for k in range(j + 1, len(normed)):
                if normed[k].startswith("tg "):
                    return i, j, k
    return -1, -1, -1


def looks_like_sev_le_cm_target_sheet(rows2d):
    header_row_idx, _, _ = find_sev_le_cm_target_block(rows2d)
    return header_row_idx >= 0


def aggregate_sev_le_cm_target_rows(rows2d, cm_to_target, stats):
    h_idx, cm_col, target_col = find_sev_le_cm_target_block(rows2d)
    if h_idx < 0:
        return
    for r in range(h_idx + 1, len(rows2d)):
        row = rows2d[r]
        if not row:
            continue
        name = get(row, cm_col)
        if name is None or str(name).strip() == "":
            continue
        name_trimmed = str(name).strip()
        if re.search(r"grand total", name_trimmed, re.I):
            continue  # pivot's own summary row, not a CM
        target = num(get(row, target_col))
        if target > 0:
            # No M1 pivot exists in this source (confirmed: just one "TG <month>" column,
            # unlike the "TG CM" sheet's two side-by-side pivots) -- m1 stays 0/unsupported
            # for any month only this source covers, same "no data" convention as elsewhere.
            cm_to_target[name_trimmed] = {"ga": target, "m1": 0}
            stats["rows"] += 1


def aggregate_sev_cm_target_rows(rows2d, cm_to_target, stats):
    """The 'TG CM' sheet lays out TWO side-by-side pivot tables on one tab (GA Mass/Migrant/
    Tourist/All on the left, M1 Mass/Migrant/Tourist/All on the right), each with its own
    'Row Labels' column -- confirmed against the real file (Target 7-eleven Aug'26_BMA
    West.xlsx): 'M1 All' sits right there with real per-CM committed M1 figures, previously
    just never read. Scanned the same way the GA column's own name column is found (nearest
    'Row Labels' to its left), independently, since a CM missing from one pivot shouldn't
    drop its value from the other."""
    h_idx = find_header_row(rows2d, ["row labels", "ga all"], 10)
    if h_idx < 0:
        return
    header = rows2d[h_idx]
    ga_all_idx = next((i for i, h in enumerate(header) if norm_header(h) == "ga all"), -1)
    if ga_all_idx < 0:
        return
    name_idx = -1
    for j in range(ga_all_idx, -1, -1):
        if norm_header(get(header, j)) == "row labels":
            name_idx = j
            break
    if name_idx < 0:
        return
    m1_all_idx = next((i for i, h in enumerate(header) if norm_header(h) == "m1 all"), -1)
    m1_name_idx = -1
    if m1_all_idx >= 0:
        for j in range(m1_all_idx, -1, -1):
            if norm_header(get(header, j)) == "row labels":
                m1_name_idx = j
                break

    for r in range(h_idx + 1, len(rows2d)):
        row = rows2d[r]
        if not row:
            continue
        name = get(row, name_idx)
        if name is None or str(name).strip() == "":
            continue
        name_trimmed = str(name).strip()
        if re.search(r"grand total", name_trimmed, re.I):
            continue  # pivot's own summary row, not a CM
        target = num(get(row, ga_all_idx))
        if target > 0:
            rec = cm_to_target.setdefault(name_trimmed, {"ga": 0, "m1": 0})
            rec["ga"] = target
            stats["rows"] += 1

    if m1_all_idx >= 0 and m1_name_idx >= 0:
        for r in range(h_idx + 1, len(rows2d)):
            row = rows2d[r]
            if not row:
                continue
            name = get(row, m1_name_idx)
            if name is None or str(name).strip() == "":
                continue
            name_trimmed = str(name).strip()
            if re.search(r"grand total", name_trimmed, re.I):
                continue
            m1_target = num(get(row, m1_all_idx))
            if m1_target > 0:
                rec = cm_to_target.setdefault(name_trimmed, {"ga": 0, "m1": 0})
                rec["m1"] = m1_target


def looks_like_sev_cm_ga_summary_sheet(rows2d):
    """An ad hoc BI export's own per-CM GA summary (seen as the 'CM' sheet in a file named
    e.g. 'Data Synergy Oct26_New.xlsx' -- naming varies month to month, so detected by
    header shape, not by file/sheet name). Two-row header: group labels (GA/Overall/Mass/
    Migrant/MM/Tourist) then leaf columns (TDS_Regions/CM/No.Store/TG/Actual/...). Matched
    on the leaf row alone -- requires CM, TG and Actual together, and deliberately NO COM
    column, since the otherwise-identical COM-level export has one and shifts every column
    one to the right (must be read with COM-aware offsets, never these fixed CM-sheet ones)."""
    for row in rows2d[:6]:
        if not row:
            continue
        normed = [norm_header(c) for c in row]
        if "cm" in normed and "com" not in normed and "tg" in normed and "actual" in normed:
            return True
    return False


def aggregate_sev_cm_ga_summary_rows(rows2d, cm_to_target, stats):
    """Reads the sheet looks_like_sev_cm_ga_summary_sheet matched. This sheet actually
    holds TWO stacked blocks sharing the exact same column layout -- a GA one first, an M1
    one afterward (confirmed against the real file: a marker cell reading "GA" sits one
    row above the first block's own header, "M1" one row above the second's, both in the
    same column as "TDS_Regions"). Reading every header-shaped row found (not just the
    first) and stopping each block's data at the NEXT one found is what makes both real --
    a single data loop running to the end of the sheet would walk straight from the GA
    block into the M1 block using the GA block's own column offsets, silently overwriting
    every CM's real GA target with its M1 figure instead (caught before this shipped).

    Only ever used to FILL a gap a CM the authoritative CM-target sheets don't have yet
    (e.g. a brand new CM from a recent reorg) -- see process_sev_files' merge order, which
    never lets this overwrite a CM those other, more authoritative sources already cover."""
    header_rows = []
    for i, row in enumerate(rows2d):
        if not row:
            continue
        normed = [norm_header(c) for c in row]
        if "cm" in normed and "com" not in normed and "tg" in normed and "actual" in normed:
            header_rows.append(i)
    if not header_rows:
        return
    for block_i, header_idx in enumerate(header_rows):
        header = rows2d[header_idx]
        # "TG"/"Actual" each repeat 5x in this row -- once per Overall/Mass/Migrant/MM/
        # Tourist sub-block -- so a {header: col} dict here would silently keep only the
        # LAST one (Tourist) instead of the Overall block's own, which is the one actually
        # wanted (caught before this shipped: it was reading each CM's Tourist-only TG as
        # if it were their real target). First occurrence, left to right, every time.
        def first_col(name, header=header):
            return next((j for j, c in enumerate(header) if norm_header(c) == name), None)
        cluster_col = first_col("tds_regions")
        cm_col = first_col("cm")
        tg_col = first_col("tg")
        if cm_col is None or tg_col is None:
            continue
        marker_row = header_idx - 1
        marker = None
        if marker_row >= 0 and rows2d[marker_row]:
            marker = norm_header(get(rows2d[marker_row], cluster_col if cluster_col is not None else 0))
        is_m1 = marker == "m1"
        block_end = header_rows[block_i + 1] if block_i + 1 < len(header_rows) else len(rows2d)
        for r in range(header_idx + 1, block_end):
            row = rows2d[r]
            if not row:
                continue
            if cluster_col is not None and not in_target_cluster(get(row, cluster_col)):
                continue
            name = get(row, cm_col)
            if name is None or str(name).strip() == "":
                continue
            name_trimmed = str(name).strip()
            if re.search(r"grand total|total", name_trimmed, re.I):
                continue  # pivot's own summary row, not a CM
            val = num(get(row, tg_col))
            if val > 0:
                rec = cm_to_target.setdefault(name_trimmed, {"ga": 0, "m1": 0})
                rec["m1" if is_m1 else "ga"] = val
                stats["rows"] += 1


ZERO_SEV_MONTH = {
    "ga": 0, "revAmt": 0, "ap100": 0, "ap49": 0, "ap99": 0, "ap199": 0, "apOver199": 0,
    "mass": 0, "migrant": 0, "tourist": 0, "dtacGa": 0, "trueGa": 0, "simTypes": {},
}


def build_sev_month_tables(month_key, all_stores, cm_targets_by_month=None):
    """Output contract -- exact column order ported from buildSevMonthTables. cm_targets_by_month
    is keyed by month (same convention as the per-store target column below it, and as
    RTR's own cmtargets sub-table) so two independent CM-target sources covering different
    months can coexist -- a month neither source covers gets an empty cmtargets table, same
    as readSevMonthGroup already tolerates via its own missing-tab fallback.

    A store with no byMonth[month_key] entry (no sales rows that month) publishes a
    zero-GA row instead of being skipped -- the exact same "dormant store vanishes"
    bug fixed on the RTR side via fill_dormant_roster_stores, just fixed here directly
    since all_stores (built across the whole loaded window) already IS this store's own
    roster: any store that sold something in ANY loaded month is a real, known store,
    so a month it sold nothing in should show 0, not disappear -- and disappearing was
    silently dropping that store's own real committed target too, not just its GA."""
    stores_rows, daily_rows, sim_rows = [], [], []
    for s in all_stores:
        bm = s["byMonth"].get(month_key) or ZERO_SEV_MONTH
        target_val = s["targetByMonth"].get(month_key, 0)
        stores_rows.append([
            s["code"], s["name"], s["cm"], s["ae"], s["khet"], s["fyai"], s["gm"], s["district"],
            bm["ga"], bm["revAmt"], bm["ap100"], bm["ap49"], bm["ap99"], bm["ap199"], bm["apOver199"],
            bm["mass"], bm["migrant"], bm["tourist"], bm["dtacGa"], bm["trueGa"], target_val,
        ])
        for st, v in bm["simTypes"].items():
            sim_rows.append([s["code"], st, v])
        for dk in sorted(k for k in s["daily"].keys() if str(k)[:6] == month_key):
            d = s["daily"][dk]
            daily_rows.append([s["code"], dk, d["ga"], d["revAmt"], d["dtacGa"], d["trueGa"],
                                d["ap49"], d["ap99"], d["ap199"], d["apOver199"],
                                d["mass"], d["migrant"], d["tourist"]])

    return {
        "stores": {
            "headers": ["code", "name", "cm", "ae", "khet", "fyai", "gm", "district", "ga", "revAmt",
                        "ap100", "ap49", "ap99", "ap199", "apOver199", "mass", "migrant", "tourist",
                        "dtacGa", "trueGa", "target"],
            "rows": stores_rows,
        },
        "daily": {
            "headers": ["code", "date", "ga", "revAmt", "dtacGa", "trueGa", "ap49", "ap99", "ap199",
                        "apOver199", "mass", "migrant", "tourist"],
            "rows": daily_rows,
        },
        "simtypes": {"headers": ["code", "simType", "ga"], "rows": sim_rows},
        "cmtargets": {
            "headers": ["cm", "target", "targetM1"],
            "rows": sorted(
                [cm, v.get("ga", 0), v.get("m1", 0)]
                for cm, v in (cm_targets_by_month or {}).get(month_key, {}).items()
            ),
        },
    }


def process_sev_files(paths):
    print(f"\nReading {len(paths)} 7-11 file(s): {', '.join(p.name for p in paths)}")
    stats = {"totalRows": 0, "clusterRows": 0}
    code_to_cm, roster_stats = {}, {"rows": 0}
    code_to_target, target_stats = {}, {"rows": 0}
    # Separate from code_to_target -- a genuinely different source (a Monitor Store-shaped
    # sheet, see aggregate_sev_monitor_store_targets) covering its OWN month, which can be
    # a different month than code_to_target's own source entirely. Kept apart so applying
    # one can never silently overwrite the other's month.
    monitor_store_targets, monitor_store_stats = {}, {"rows": 0}
    # Another independent source, same reasoning -- an "LE by store" file's per-store
    # September target (see aggregate_sev_le_by_store_targets), applied AFTER
    # monitor_store_targets so it wins on any overlapping code, since it's the more
    # complete/authoritative of the two (Monitor Store's own target columns turned out to
    # be entirely blank once checked across every row, not just coincidentally zero).
    le_by_store_targets, le_by_store_stats = {}, {"rows": 0}
    cm_to_target, cm_target_stats = {}, {"rows": 0}
    # Kept separate from cm_to_target, same reasoning as le_by_store_targets above -- the
    # PV sheet's CM-level rollup can cover a DIFFERENT month than whatever the older named
    # "TG CM"-style sheet covers, and a single flat dict can't hold two months' values for
    # the same CM name at once without one silently clobbering the other.
    le_cm_to_target, le_cm_target_stats = {}, {"rows": 0}
    # Another independent, lower-precedence source -- an ad hoc BI export's own "CM" summary
    # sheet (see looks_like_sev_cm_ga_summary_sheet), only ever used to fill in a CM the
    # authoritative cm_to_target/le_cm_to_target sheets don't have yet (e.g. a brand new CM
    # a recent reorg introduced that those files haven't caught up to).
    cm_ga_summary_to_target, cm_ga_summary_stats = {}, {"rows": 0}
    sheet_errors = []

    # ---- pass 1: cheap metadata only (sheet names, header row, and -- for anything that
    # might be a sales sheet -- a day-column-only scan) to figure out which months are
    # even worth fully reading, before paying for the actually expensive part below.
    # Deliberately NOT based on sheet/file NAMING ("Q1_Database" / "Jan" implies Jan-Mar,
    # but isn't guaranteed to only CONTAIN Jan-Mar rows) -- it's grounded in the same
    # TM_KEY_DAY column the real aggregation uses for real, just scanned alone instead of
    # the whole row. A sheet whose day-scan shows even one row in the window this run will
    # keep is never skipped, so this can't drop anything the real aggregation would have kept.
    sheet_meta = []  # {fileName, name, isOther, monthsTouched}
    for path in paths:
        print(f"  scanning {path.name}...", end="", flush=True)
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            sheet_errors.append(f"{path.name}: could not open ({e})")
            print(" failed to open")
            continue
        for name in wb.sheetnames:
            ws = wb[name]
            header_row = next(ws.iter_rows(max_row=1, values_only=True), None)
            header = list(header_row) if header_row else []
            # looks_like_sev_cm_ga_summary_sheet's header is two rows deep (group labels,
            # then leaf columns) -- the other is_other checks here only ever need row 1.
            second_row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True), None)
            second_row = list(second_row) if second_row else []
            is_other = (
                looks_like_sev_cm_target_sheet(name)
                or looks_like_sev_target_sheet(name)
                or looks_like_roster_sheet([header])
                or looks_like_sev_le_by_store_sheet([header])
                or looks_like_sev_le_cm_target_sheet([header])
                or looks_like_sev_cm_ga_summary_sheet([header, second_row])
            )
            months_touched = None
            if not is_other:
                idx = {str(h).strip() if h is not None else "": i for i, h in enumerate(header)}
                day_col = idx.get("TM_KEY_DAY")
                if day_col is not None:
                    months_touched = sev_months_touched(ws, day_col)
            sheet_meta.append({"fileName": path.name, "name": name, "isOther": is_other,
                                "monthsTouched": months_touched})
        wb.close()
        print(" done")

    # The window this run will actually keep -- newest KEEP_MONTHS distinct months found
    # across every sales sheet's day-scan. A sheet with NO rows in this window can't
    # contribute anything build_payload's own retention trim wouldn't just discard anyway,
    # so fully reading it would be pure waste. isOther sheets (target/CM-target/roster,
    # always needed regardless of month) and ones whose months couldn't be determined
    # (missing TM_KEY_DAY -- let aggregate_sev_rows raise its own clear error for that,
    # rather than silently skip) are never skipped, only real sales sheets are.
    all_months_touched = set()
    for sh in sheet_meta:
        if sh["monthsTouched"]:
            all_months_touched |= sh["monthsTouched"]
    local_window = set(sorted(all_months_touched)[-KEEP_MONTHS:]) if all_months_touched else None

    def sheet_needed(sh):
        return (
            sh["isOther"]
            or sh["monthsTouched"] is None
            or local_window is None
            or bool(sh["monthsTouched"] & local_window)
        )

    all_sheets = []  # {fileName, name, rows2D}
    skipped_sheets = []
    for path in paths:
        this_file_meta = [sh for sh in sheet_meta if sh["fileName"] == path.name]
        names_needed = [sh["name"] for sh in this_file_meta if sheet_needed(sh)]
        names_skipped = [sh["name"] for sh in this_file_meta if not sheet_needed(sh)]
        skipped_sheets.extend(f"{path.name} / {n}" for n in names_skipped)
        if not names_needed:
            print(f"  {path.name}: no sheets in the {KEEP_MONTHS}-month window, skipped entirely")
            continue
        # Same reasoning as process_rtr_workbook's per-sheet print -- one of these files
        # (the store database especially) is itself hundreds of thousands of rows, so this
        # prints AS EACH FILE finishes instead of going silent until all 6 are done.
        print(f"  reading {path.name}...", end="", flush=True)
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            sheet_errors.append(f"{path.name}: could not open ({e})")
            print(" failed to open")
            continue
        for name in names_needed:
            all_sheets.append({"fileName": path.name, "name": name, "rows2D": sheet_to_rows(wb[name])})
        wb.close()
        print(f" done ({len(names_skipped)} sheet(s) skipped)" if names_skipped else " done")

    if skipped_sheets:
        print(f"  skipped {len(skipped_sheets)} sheet(s) with no rows in the {KEEP_MONTHS}-month "
              f"window this run will keep: {skipped_sheets}")

    sales_sheets, other_sheets = [], []
    for sh in all_sheets:
        if (looks_like_sev_cm_target_sheet(sh["name"]) or looks_like_sev_target_sheet(sh["name"])
                or looks_like_roster_sheet(sh["rows2D"]) or looks_like_sev_monitor_store_sheet(sh["rows2D"])
                or looks_like_sev_le_by_store_sheet(sh["rows2D"]) or looks_like_sev_le_cm_target_sheet(sh["rows2D"])
                or looks_like_sev_cm_ga_summary_sheet(sh["rows2D"])):
            other_sheets.append(sh)
        else:
            sh["dayCounts"] = sev_count_rows_by_day(sh["rows2D"])
            sales_sheets.append(sh)

    day_owner = {}
    for idx, sh in enumerate(sales_sheets):
        for day, count in sh["dayCounts"].items():
            cur = day_owner.get(day)
            if cur is None or count > sales_sheets[cur]["dayCounts"][day]:
                day_owner[day] = idx

    stores_map = {}
    for idx, sh in enumerate(sales_sheets):
        # The real bulk of this script's time is here, not the file reads above -- millions of
        # rows across every sales sheet get scanned per sheet. Same "print as each one finishes"
        # treatment so a run doesn't look stalled for the minute or more this can take.
        print(f"  aggregating {sh['fileName']} / {sh['name']}...", end="", flush=True)
        try:
            aggregate_sev_rows(sh["rows2D"], stores_map, stats, day_owner, idx)
            print(" done")
        except Exception as e:
            print(" failed")
            sheet_errors.append(f"{sh['fileName']} / {sh['name']}: {e}")

    # Which file each target actually came from -- its own filename ("Target 7-eleven
    # Aug'26_BMA West.xlsx" / "AE Monitoring_Sep'26.xlsx") is the only place that month is
    # recorded at all; the sheets inside carry no month of their own. First file to
    # contribute any rows wins PER SOURCE -- code_to_target and monitor_store_targets are
    # independent sources that can (and currently do) each cover a DIFFERENT month, not one
    # target file overall.
    target_file_name = None
    cm_target_file_name = None
    monitor_store_file_name = None
    le_by_store_file_name = None
    le_cm_target_file_name = None
    cm_ga_summary_file_name = None
    for sh in other_sheets:
        try:
            if looks_like_sev_cm_ga_summary_sheet(sh["rows2D"]):
                before = len(cm_ga_summary_to_target)
                aggregate_sev_cm_ga_summary_rows(sh["rows2D"], cm_ga_summary_to_target, cm_ga_summary_stats)
                if cm_ga_summary_file_name is None and len(cm_ga_summary_to_target) > before:
                    cm_ga_summary_file_name = sh["fileName"]
            elif looks_like_sev_cm_target_sheet(sh["name"]):
                before = len(cm_to_target)
                aggregate_sev_cm_target_rows(sh["rows2D"], cm_to_target, cm_target_stats)
                if cm_target_file_name is None and len(cm_to_target) > before:
                    cm_target_file_name = sh["fileName"]
            elif looks_like_sev_target_sheet(sh["name"]):
                before = len(code_to_target)
                aggregate_sev_target_rows(sh["rows2D"], code_to_target, target_stats)
                if target_file_name is None and len(code_to_target) > before:
                    target_file_name = sh["fileName"]
            elif looks_like_sev_monitor_store_sheet(sh["rows2D"]):
                before = len(monitor_store_targets)
                aggregate_sev_monitor_store_targets(sh["rows2D"], monitor_store_targets, monitor_store_stats)
                if monitor_store_file_name is None and len(monitor_store_targets) > before:
                    monitor_store_file_name = sh["fileName"]
            elif looks_like_sev_le_cm_target_sheet(sh["rows2D"]):
                before = len(le_cm_to_target)
                aggregate_sev_le_cm_target_rows(sh["rows2D"], le_cm_to_target, le_cm_target_stats)
                if le_cm_target_file_name is None and len(le_cm_to_target) > before:
                    le_cm_target_file_name = sh["fileName"]
            elif looks_like_sev_le_by_store_sheet(sh["rows2D"]):
                # This file's month (e.g. "Sep 2026") comes from its own filename, same
                # convention as target_file_name/monitor_store_file_name -- converted to
                # the 3-letter abbreviation aggregate_sev_le_by_store_targets keys its
                # columns by (see find_sev_le_by_ae_columns).
                month_key = parse_month_from_filename(sh["fileName"])
                month_abbr = MONTH_LABEL[int(month_key[4:6]) - 1].lower() if month_key else None
                if month_abbr:
                    before = len(le_by_store_targets)
                    aggregate_sev_le_by_store_targets(sh["rows2D"], le_by_store_targets,
                                                       le_by_store_stats, month_abbr)
                    if le_by_store_file_name is None and len(le_by_store_targets) > before:
                        le_by_store_file_name = sh["fileName"]
            else:
                aggregate_roster_rows(sh["rows2D"], code_to_cm, roster_stats)
        except Exception as e:
            sheet_errors.append(f"{sh['fileName']} / {sh['name']}: {e}")

    if sheet_errors:
        print("  sheet warnings (non-fatal):")
        for e in sheet_errors:
            print(f"    - {e}")

    if not stores_map:
        raise SystemExit(
            f"No stores found in the target cluster across {len(paths)} file(s) "
            f"({stats['totalRows']} total rows scanned)."
        )

    all_stores = list(stores_map.values())
    # Precedence flipped from the browser's original (roster always wins): starting with the
    # new monthly "AE Monitoring" format, each month's own Data sheet carries CM directly per
    # row (latest month wins -- see cmMonthKey above), and that's now the authoritative source.
    # The old static roster file is only a fallback for a store no loaded month's Data sheet
    # ever assigned a CM to at all.
    #
    # BUT: verified against a real file, one store (71117324) came back with a CM name
    # ("นุกูล พรมศร") that doesn't belong to this territory's roster at all. Confirmed with the
    # user this means the store has actually moved OUT of this territory (a different CM's own
    # AE, "นรเศรษฐ์ นุชสา", consistently reports there across every row that month) -- not that
    # the value is a typo to paper over with a stale roster guess. So an unrecognized CM name
    # drops the store the same way a wholly-unmapped one already is (see unmapped_count below),
    # rather than silently keeping it under whatever the old roster happened to say. Roster is
    # still the right fallback for the OTHER case -- a store some month's Data sheet simply
    # never carried a CM for at all (s["cm"] is None, not an unrecognized value) -- that's a
    # genuine gap to fill in, not a contradicting signal.
    #
    # known_cms must be scoped to CMs actually seen managing IN-CLUSTER stores, not
    # code_to_cm's full value set -- the roster file is company-wide (13 distinct CMs across
    # every territory, confirmed against the real file), not specific to this cluster,  and
    # "นุกูล พรมศร" is a real CM in it -- just for stores elsewhere in the company, which is
    # exactly why his name here was the signal that store had left this territory, not proof
    # he's a legitimate CM for it.
    in_cluster_codes = set(stores_map.keys())
    known_cms = {cm for code, cm in code_to_cm.items() if code in in_cluster_codes}
    # Company-wide (not cluster-scoped): a per-row CM name the roster recognizes for SOME
    # other territory is "นุกูล พรมศร" below. A name the roster doesn't recognize ANYWHERE,
    # though, isn't that signal at all -- it's a CM the roster file simply hasn't caught up
    # to yet (e.g. a just-announced reorg splitting one CM's book into two new CMs). That
    # case should keep the fresh per-row value, not discard it to UNMAPPED_CM.
    company_wide_cms = set(code_to_cm.values())
    for s in all_stores:
        if s["cm"] is None:
            s["cm"] = code_to_cm.get(s["code"]) or UNMAPPED_CM
        elif s["cm"] not in known_cms and s["cm"] in company_wide_cms:
            s["cm"] = UNMAPPED_CM

    # A store's CM comes from its own latest sales row, so a store with no sale yet in the
    # newest month still carries the CM from before a reorg (e.g. 18 of one AE's stores still
    # showing the old CM while the AE's other stores show the new one). Each AE's current CM
    # is the majority CM among their stores that DID sell in the newest month; their other
    # stores follow it, so one AE never appears under two CMs at once.
    newest_mk = max((s["cmMonthKey"] for s in all_stores if s.get("cmMonthKey")), default=None)
    ae_votes = {}
    for s in all_stores:
        if (s.get("cmMonthKey") == newest_mk and s["cm"] != UNMAPPED_CM
                and s["ae"] and s["ae"] != UNMAPPED_AE):
            ae_votes.setdefault(s["ae"], Counter())[s["cm"]] += 1
    ae_current_cm = {ae: votes.most_common(1)[0][0] for ae, votes in ae_votes.items()}
    followed = 0
    for s in all_stores:
        new_cm = ae_current_cm.get(s["ae"])
        if new_cm and s.get("cmMonthKey") != newest_mk and s["cm"] != UNMAPPED_CM and s["cm"] != new_cm:
            s["cm"] = new_cm
            followed += 1
    if followed:
        print(f"  {followed} store(s) with no {newest_mk} sale yet moved to their AE's current CM")

    unmapped_count = unmapped_ga = 0
    stores = []
    for s in all_stores:
        if s["cm"] == UNMAPPED_CM:
            unmapped_count += 1
            unmapped_ga += s["ga"]
            continue
        stores.append(s)

    months_available = sorted({mk for s in stores for mk in s["byMonth"].keys()})
    # The target file's OWN month (parsed from its filename), not just "whichever loaded
    # month happens to be newest" -- that assumption broke the moment a newer month's sales
    # data (e.g. this month's own "AE Monitoring" file) started arriving in a separate,
    # later file than the target/commit file itself. Falls back to the old newest-month
    # guess only if the filename couldn't be parsed at all.
    target_month_key = (
        (target_file_name and parse_month_from_filename(target_file_name))
        or (months_available[-1] if months_available else None)
    )
    if code_to_target and target_month_key:
        for s in stores:
            t = code_to_target.get(s["code"])
            if t is not None:
                s["targetByMonth"][target_month_key] = t

    # Independent of target_month_key above -- Monitor Store's own file covers its own
    # month (currently September, via AE Monitoring_Sep'26.xlsx), which can be a totally
    # different month than code_to_target's source. Both get applied into the same
    # targetByMonth dict, keyed by their own month, so neither can overwrite the other.
    monitor_store_month_key = (
        monitor_store_file_name and parse_month_from_filename(monitor_store_file_name)
    )
    if monitor_store_targets and monitor_store_month_key:
        for s in stores:
            t = monitor_store_targets.get(s["code"])
            if t is not None:
                s["targetByMonth"][monitor_store_month_key] = t

    # Applied last so it wins over monitor_store_targets on any overlapping code -- see the
    # comment where le_by_store_targets is initialized above.
    le_by_store_month_key = (
        le_by_store_file_name and parse_month_from_filename(le_by_store_file_name)
    )
    if le_by_store_targets and le_by_store_month_key:
        for s in stores:
            t = le_by_store_targets.get(s["code"])
            if t is not None:
                s["targetByMonth"][le_by_store_month_key] = t

    cm_target_month_key = (
        (cm_target_file_name and parse_month_from_filename(cm_target_file_name))
        or target_month_key
    )
    le_cm_target_month_key = (
        le_cm_target_file_name and parse_month_from_filename(le_cm_target_file_name)
    )
    # No month of its own in the sheet (unlike cm_to_target/le_cm_to_target's source files,
    # which carry an explicit "TG <month>" column) -- same fallback as target_month_key.
    cm_ga_summary_month_key = (
        (cm_ga_summary_file_name and parse_month_from_filename(cm_ga_summary_file_name))
        or (months_available[-1] if months_available else None)
    )

    # Merged by month, not into one flat dict -- see le_cm_to_target's own comment above.
    # Each source only ever writes into its OWN month's slot, so an overlapping CM name
    # across two sources can only collide if they both genuinely claim the same month
    # (a real conflict worth letting one win), never across different months.
    cm_target_by_month = {}
    if cm_to_target and cm_target_month_key:
        cm_target_by_month.setdefault(cm_target_month_key, {}).update(cm_to_target)
    if le_cm_to_target and le_cm_target_month_key:
        cm_target_by_month.setdefault(le_cm_target_month_key, {}).update(le_cm_to_target)
    # Lowest precedence, gap-fill only -- never overwrites a CM cm_to_target/le_cm_to_target
    # already gave a target for that same month, only adds one for a CM neither of those
    # authoritative sources has yet (see cm_ga_summary_to_target's own comment above).
    filled_cms = []
    if cm_ga_summary_to_target and cm_ga_summary_month_key:
        month_targets = cm_target_by_month.setdefault(cm_ga_summary_month_key, {})
        for cm, v in cm_ga_summary_to_target.items():
            if cm not in month_targets:
                month_targets[cm] = v
                filled_cms.append(cm)

    print(f"  {stats['totalRows']} rows scanned, {stats['clusterRows']} in cluster, "
          f"{len(stores)} stores mapped ({unmapped_count} unmapped / GA {unmapped_ga:.0f} dropped)")
    print(f"  roster rows: {roster_stats['rows']}, store targets: {target_stats['rows']} "
          f"(month: {target_month_key}), Monitor Store targets: {monitor_store_stats['rows']} "
          f"(month: {monitor_store_month_key}), LE by store targets: {le_by_store_stats['rows']} "
          f"(month: {le_by_store_month_key}), CM targets: {cm_target_stats['rows']} "
          f"(month: {cm_target_month_key}), LE CM targets: {le_cm_target_stats['rows']} "
          f"(month: {le_cm_target_month_key})")
    if cm_ga_summary_to_target:
        print(f"  CM GA summary sheet: {cm_ga_summary_stats['rows']} CM(s) read "
              f"(month: {cm_ga_summary_month_key}); filled in a missing target for: "
              f"{sorted(filled_cms) if filled_cms else '(none -- already covered)'}")
    print(f"  months available: {months_available}")
    for mk in months_available:
        total_ga = sum(s["byMonth"][mk]["ga"] for s in stores if mk in s["byMonth"])
        print(f"    {mk}: GA {total_ga:.0f}")

    return stores, months_available, cm_target_by_month


# ═════════════════════════════════════════════════════════════════════════
# PUBLISH -- build the payload, apply retention, POST to the Apps Script
# ═════════════════════════════════════════════════════════════════════════

def build_payload(rtr_months, sev_stores, sev_months_available, cm_targets_by_month=None,
                   existing_tables=None, existing_meta=None):
    """Builds the publish payload, merging this run's freshly-computed local months with
    whatever the Sheet already has (existing_tables/existing_meta, from fetch_sync_data) --
    local data always wins for a month it actually covers ("is there something new for this
    month" is exactly "did this run's local files produce it"); a month the Sheet has that
    this run's local files don't cover is carried through UNCHANGED rather than silently
    dropped just because a source file happened to be missing this run. The KEEP_MONTHS cap
    is applied to the MERGED set, not just what was freshly computed, so the Sheet always
    ends up holding exactly the newest KEEP_MONTHS of each feed -- a genuinely new month
    pushes the oldest one off for real (see deleteTabs below), not just out of what the
    dashboard bothers to look at."""
    existing_tables = existing_tables or {}
    existing_meta = existing_meta or {}

    tables = {}
    rtr_keys, sev_keys = [], []

    for m in rtr_months:
        d = parse_sheet_date(m["sheetName"])
        if not d:
            continue
        # Same fallback as the browser's own syncBuildPayloads: a sheet name with no year
        # (e.g. "Apr" rather than "Apr.26") falls back to the CURRENT real year here, not
        # the majority-vote inferred year sortMonthsChronologically used for ordering --
        # those are deliberately two different fallbacks in the original JS.
        year = d["year"] or date.today().year
        key = f"rtr-{year}-{d['monthIndex'] + 1:02d}"
        rtr_keys.append(key)
        for suffix, table in build_rtr_month_tables(m).items():
            tables[f"{key}-{suffix}"] = table

    for mk in sev_months_available:
        key = f"sev-{mk[:4]}-{mk[4:6]}"
        sev_keys.append(key)
        for suffix, table in build_sev_month_tables(mk, sev_stores, cm_targets_by_month).items():
            tables[f"{key}-{suffix}"] = table

    def group_key_of(tab):
        # tab looks like "rtr-2026-08-stores" -> group key "rtr-2026-08"
        parts = tab.split("-")
        return "-".join(parts[:3]) if len(parts) >= 3 else tab

    # Retention, now applied to the UNION of what the Sheet already had and what this run
    # actually found locally -- matching SYNC.KEEP_MONTHS / syncPrune's month-key sort
    # (lexicographic "YYYY-MM" == chronological). A month present in both sets uses the
    # fresh local version rebuilt into `tables` above, never the stale existing one.
    existing_rtr_keys = set(existing_meta.get("rtr_months") or [])
    existing_sev_keys = set(existing_meta.get("sev_months") or [])

    def merge_and_trim(fresh_keys, existing_keys):
        return sorted(set(fresh_keys) | existing_keys)[-KEEP_MONTHS:]

    kept_rtr = merge_and_trim(rtr_keys, existing_rtr_keys)
    kept_sev = merge_and_trim(sev_keys, existing_sev_keys)
    kept = set(kept_rtr) | set(kept_sev)

    final_tables = {}
    for tab, table in tables.items():
        if group_key_of(tab) in kept:
            final_tables[tab] = table
    # Months the Sheet already had that this run didn't recompute (no local file covered
    # them this time) but that still survive the cap -- carried through byte-for-byte from
    # what fetch_sync_data just read, never touched or reprocessed. This is exactly what
    # stops a temporarily-missing source file from making an old month vanish.
    for tab, table in existing_tables.items():
        if tab in final_tables:
            continue  # fresh local version already wins for this tab
        if group_key_of(tab) in kept:
            final_tables[tab] = table

    # Now a REAL cap (by request): anything the Sheet has whose month fell out of the
    # merged+trimmed set gets explicitly deleted, so the Sheet always ends up holding
    # exactly KEEP_MONTHS months of each feed -- not the previous, more conservative
    # default of just leaving old tabs stop being refreshed.
    delete_tabs = sorted(tab for tab in existing_tables if group_key_of(tab) not in kept)

    return {
        "tables": final_tables,
        "meta": {
            "rtr_months": kept_rtr,
            "sev_months": kept_sev,
            "updated_at": int(time.time() * 1000),
        },
        "deleteTabs": delete_tabs,
    }


def post_to_sync(action, secret, extra=None, timeout=120, retries=3, backoff=4):
    """POSTs one action to the Apps Script sync endpoint, retrying a TRANSIENT failure
    (a network error, an HTTP error, or a response that didn't parse as JSON -- e.g. Apps
    Script cold-starting right after a redeploy, or a bare blip like the one HTTP 404
    observed in practice that a bare retry moments later cleared on its own) with a short
    backoff between attempts. process_rtr_workbook + process_sev_files together can take a
    minute or more (process_sev_files alone scans several million rows), and both already
    ran by the time this is ever called -- so silently retrying a POST here is far cheaper
    than forcing a full rerun from scratch just because a network blip happened to land on
    this one request. Does NOT retry an application-level failure (result.get("ok") is
    False) -- that's each caller's own decision to raise, since retrying an invalid secret
    or a rejected payload would just fail the exact same way every time, burning through
    retries for nothing. Raises SystemExit with the last error if every attempt fails."""
    body = json.dumps({"action": action, "secret": secret, **(extra or {})}).encode("utf-8")
    last_err = None
    for attempt in range(1, retries + 1):
        req = urllib.request.Request(
            SYNC_URL, data=body, method="POST",
            headers={"Content-Type": "text/plain;charset=utf-8"},
        )
        try:
            with urllib.request.urlopen(req, timeout=timeout) as res:
                raw = res.read().decode("utf-8")
            return json.loads(raw)
        except urllib.error.HTTPError as e:
            last_err = f"HTTP {e.code}\n{e.read().decode('utf-8', 'replace')[:500]}"
        except urllib.error.URLError as e:
            last_err = str(e.reason)
        except OSError as e:
            # socket.timeout while READING the body (after the connection already opened)
            # isn't wrapped in URLError -- without this a slow Apps Script reply crashed the
            # run with a traceback instead of retrying.
            last_err = f"network error: {e or type(e).__name__}"
        except json.JSONDecodeError:
            # A handful of tables (54+ tabs' worth) can push this response into multi-MB
            # territory -- if Apps Script hit its own 6-minute execution cap or just needed
            # a moment to spin up, the response can come back empty or truncated instead of
            # a clean JSON error.
            last_err = f"response wasn't JSON (got {len(raw)} bytes)\nFirst 300 chars:\n{raw[:300]}"
        if attempt < retries:
            wait = backoff * attempt
            print(f"  {action} attempt {attempt}/{retries} failed ({last_err.splitlines()[0]}) "
                  f"-- retrying in {wait}s...")
            time.sleep(wait)
    raise SystemExit(f"{action} failed after {retries} attempts: {last_err}")


def fetch_sync_data(secret):
    """Reads the Sheet's current tables + meta (see getSyncData_ in the Apps Script) so
    build_payload can merge fresh local months into it instead of blindly overwriting.
    Raises on any failure -- including HTTP/network errors -- rather than treating a failed
    fetch the same as "nothing synced yet", which would defeat the whole point: a merge that
    silently proceeds as if the Sheet were empty is exactly the data-loss path this exists
    to avoid."""
    result = post_to_sync("getSyncData", secret, timeout=180)
    if not result.get("ok"):
        raise SystemExit(f"getSyncData failed: {result.get('error')}")
    return result


def publish(payload, secret):
    print(f"\nUploading {len(payload['tables'])} tabs to the Sheet...")
    result = post_to_sync("publishTerritory", secret, extra=payload, timeout=120)
    if not result.get("ok"):
        raise SystemExit(f"Upload failed: {result.get('error')}")
    print(f"Done -- {result['tablesWritten']} tables written.")


# ═════════════════════════════════════════════════════════════════════════
# PIPELINE CACHE -- most of a run's time is spent re-scanning source files
# whose content hasn't actually changed since the last run (Q2/Q3_Database,
# Target Aug, the roster file, LE-by-store all sit static for days at a time;
# only the RTR workbook and AE Monitoring get refreshed regularly). There's
# no cheap way to know a file's CONTENT is unchanged without reading it, but
# mtime+size is a good enough proxy in practice -- editing/replacing a file
# always changes at least one of the two, and nothing else in this workflow
# rewrites a file byte-for-byte with the same mtime and size but different
# content. Cached at the OUTPUT level (the finished rtr_months / sev_stores
# etc.), not per-file, since process_sev_files' own aggregation genuinely
# depends on every file together (day_owner picks, per calendar day, whichever
# loaded file has the most complete data for it -- a cross-file decision, not
# one any single file's cache entry could safely capture on its own). So a
# cache hit requires the WHOLE file set (paths + mtimes + sizes) to match
# exactly what produced the cached result; any single file changing (new
# content, a file added/removed) invalidates that side's cache entirely and
# falls back to a full fresh read, same as if this cache didn't exist.
# ═════════════════════════════════════════════════════════════════════════
PIPELINE_CACHE_FILE = TERRITORY_DIR / ".pipeline_cache.json"


def file_signature(path):
    st = path.stat()
    return {"path": str(path), "mtime": st.st_mtime, "size": st.st_size}


def files_signature(paths):
    return [file_signature(p) for p in paths]


# Folded into every cache signature below (see script_signature's own callers) -- a cache
# entry is only ever a hit if the SOURCE FILES are unchanged AND this script itself hasn't
# changed since. Without this, editing the script's own aggregation logic (e.g. adding a
# new field to a cmRollup dict) while the underlying xlsx files stay untouched would still
# read back the OLD cached shape, silently missing the new field -- exactly the KeyError
# this caught in practice the first time this ran after adding targetMassMigrant/
# targetTourist. Self-healing: the cache key's own SHAPE changed here too, so any
# pre-existing cache file just stops matching and falls back to a fresh read, no manual
# cache-clearing step required.
def script_signature():
    return file_signature(Path(__file__).resolve())


def load_pipeline_cache():
    try:
        return json.loads(PIPELINE_CACHE_FILE.read_text(encoding="utf-8"))
    except Exception:
        # Missing (first run ever), corrupt, or unreadable -- either way, treat it the
        # same as "nothing cached yet" rather than letting a bad cache file break a
        # real publish run. Worst case this just means a fresh read, same as before
        # this cache existed.
        return {}


def save_pipeline_cache(cache):
    try:
        PIPELINE_CACHE_FILE.write_text(
            json.dumps(cache, ensure_ascii=False), encoding="utf-8"
        )
    except Exception as e:
        # Caching is a nice-to-have speedup, not a requirement -- a write failure
        # (disk full, permissions) shouldn't fail a publish that otherwise succeeded.
        print(f"  (could not save pipeline cache, non-fatal: {e})")


PROGRESS_TOTAL_STEPS = 5


def progress(step, label, width=28):
    filled = int(width * step / PROGRESS_TOTAL_STEPS)
    bar = "█" * filled + "░" * (width - filled)
    pct = int(100 * step / PROGRESS_TOTAL_STEPS)
    print(f"\n[{bar}] {pct:3d}%  Step {step}/{PROGRESS_TOTAL_STEPS}: {label}")


def main():
    if not RTR_XLSX_PATH:
        raise SystemExit(f"No .xlsx found in {RTR_DIR}")
    if not SEV_PATHS:
        raise SystemExit(f"No .xlsx found in {SEV_DIR}")
    if not SYNC_SECRET_FILE.exists():
        raise SystemExit(
            f"Not found: {SYNC_SECRET_FILE}\n"
            "This should already exist -- if missing, generate one with "
            "`openssl rand -hex 24` and set the same value as TERRITORY_SYNC_SECRET "
            "in the Apps Script project's Script Properties."
        )
    secret = SYNC_SECRET_FILE.read_text(encoding="utf-8").strip()

    cache = load_pipeline_cache()
    script_sig = script_signature()

    progress(1, "Reading RTR workbook")
    rtr_sig = {
        "file": file_signature(RTR_XLSX_PATH),
        "script": script_sig,
        "routing": file_signature(ROUTING_XLSX_PATH) if ROUTING_XLSX_PATH else None,
    }
    if cache.get("rtr", {}).get("signature") == rtr_sig:
        print(f"\nRTR workbook unchanged since last run ({RTR_XLSX_PATH.name}) -- "
              f"using cached results, skipping re-read.")
        rtr_months = cache["rtr"]["months"]
    else:
        rtr_months = process_rtr_workbook(RTR_XLSX_PATH, routing_path=ROUTING_XLSX_PATH)
        cache["rtr"] = {"signature": rtr_sig, "months": rtr_months}
        save_pipeline_cache(cache)

    progress(2, "Reading 7-11 files")
    sev_sig = {"files": files_signature(SEV_PATHS), "script": script_sig}
    if cache.get("sev", {}).get("signature") == sev_sig:
        print(f"7-11 files unchanged since last run ({len(SEV_PATHS)} files) -- "
              f"using cached results, skipping re-read.")
        sev_stores = cache["sev"]["stores"]
        sev_months = cache["sev"]["months_available"]
        cm_target_by_month = cache["sev"]["cm_target_by_month"]
    else:
        sev_stores, sev_months, cm_target_by_month = process_sev_files(SEV_PATHS)
        cache["sev"] = {
            "signature": sev_sig, "stores": sev_stores,
            "months_available": sev_months, "cm_target_by_month": cm_target_by_month,
        }
        save_pipeline_cache(cache)

    progress(3, "Fetching the Sheet's current state")
    print("Fetching the Sheet's current state (to merge into, not overwrite)...")
    sync_data = fetch_sync_data(secret)
    existing_tables = sync_data.get("tables") or {}
    existing_meta = sync_data.get("meta") or {}
    if existing_meta:
        print(f"  Sheet currently has: rtr {existing_meta.get('rtr_months')}, "
              f"sev {existing_meta.get('sev_months')}")
    else:
        print("  nothing synced yet -- this will be the first push.")

    progress(4, "Preparing tables to publish")
    payload = build_payload(rtr_months, sev_stores, sev_months, cm_target_by_month,
                             existing_tables, existing_meta)
    print(f"Will publish: rtr {payload['meta']['rtr_months']}, sev {payload['meta']['sev_months']}")
    if payload["deleteTabs"]:
        print(f"  will delete {len(payload['deleteTabs'])} old tab(s) that fell out of the "
              f"{KEEP_MONTHS}-month window: {payload['deleteTabs']}")

    if "--dry-run" in sys.argv:
        progress(5, "Dry run -- not uploading")
        return

    progress(5, "Uploading to Google Sheets")
    publish(payload, secret)


RTR_XLSX_PATH = _find_rtr_workbook(RTR_DIR)
ROUTING_XLSX_PATH = _find_routing_workbook(RTR_DIR, exclude=RTR_XLSX_PATH)
SEV_PATHS = sorted(p for p in SEV_DIR.glob("*.xlsx") if not p.name.startswith("~$"))

if __name__ == "__main__":
    main()
